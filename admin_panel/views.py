#admin_panel/views.py
import re
import json
from django.views.decorators.http import require_http_methods
from django.shortcuts import render,redirect
from django.contrib.auth.decorators import login_required, user_passes_test
from django.views.decorators.http import require_GET, require_POST
from django.db.models import F, Q
import csv
from django.forms.models import model_to_dict
from openpyxl import styles
from trainer.models import AttendanceRecord
from django.http import JsonResponse, HttpResponse, Http404
from django.core.exceptions import FieldDoesNotExist
from django.core.exceptions import FieldDoesNotExist
from django.utils.timezone import localtime
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter
from collections import defaultdict
from practicetest.models import PracticeTest, PracticeQuestion, PracticeResult, ScheduledPracticeTest
from exam.models import Exam, ExamQuestion, ExamResult, ScheduledExam
from pre_assessment.models import PreAssessmentExam
from college.forms import CollegeForm
from college.models import College,Course,Section
from material.models import Domain,SubDomain,Module
from student.models import StudentProfile
from django.db.models import F
from django.views.decorators.http import require_http_methods
from django.core.paginator import Paginator
from django.db.models import Q
from django.db.models import Count
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image
from openpyxl.chart import BarChart, Reference, PieChart
from openpyxl.chart.label import DataLabelList
from io import BytesIO
from datetime import datetime
from django.http import HttpResponse
import json, csv
from trainer.models import TrainerProfile  # <-- missing import
from college.models import College
import csv
from django.http import JsonResponse, HttpResponse
from django.utils.timezone import localtime
from django.contrib.auth.decorators import login_required, user_passes_test
from django.utils.dateparse import parse_date
from django.forms.models import model_to_dict
from django.db.models import Avg, Count
from django.shortcuts import render, get_object_or_404
from django.forms.models import model_to_dict
from django.db import transaction, IntegrityError
from student.models import StudentProfile, StudentFeedback, SessionTrainerFeedback
from college.models import College, Course, Section
from admin_panel.models import SessionSchedule
from .models import SessionSchedule, ValidationError
from trainer.models import TrainerSessionReport
from trainer.models import TrainerProfile
from django.db.models import Prefetch
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image
from openpyxl.chart import BarChart, Reference, PieChart
from openpyxl.chart.label import DataLabelList
from io import BytesIO
from datetime import datetime
from django.contrib import messages
from .forms import (
    StudentRegistrationForm, BulkStudentUploadForm,
    TrainerRegistrationForm, TpoRegistrationForm, StudentEditForm
)

def _student_name(student):
    """Helper function to safely get student full name."""
    try:
        if hasattr(student.user, 'get_full_name'):
            name = student.user.get_full_name()
            if name and name.strip():
                return name
        first = getattr(student.user, 'first_name', '') or ''
        last = getattr(student.user, 'last_name', '') or ''
        name = f'{first} {last}'.strip()
        return name if name else getattr(student.user, 'username', '—')
    except Exception:
        return '—'


def _format_question_time_map(time_map):
    if not isinstance(time_map, dict) or not time_map:
        return "-"

    def _sort_key(item):
        key = str(item[0])
        return (0, int(key)) if key.isdigit() else (1, key)

    pairs = sorted(time_map.items(), key=_sort_key)
    return " | ".join([f"Q{k}: {v}s" for k, v in pairs])


def _submission_mode_for_exam(result):
    status_key = str(getattr(result, "submission_status", "") or "").lower()
    tab_limit = 3
    if result.scheduled_exam:
        tab_limit = int(getattr(result.scheduled_exam, "allowed_tab_switches", 3) or 3)
    tab_switch_count = getattr(result, "tab_switch_count", 0) or 0

    if tab_switch_count >= tab_limit or status_key == "accidental_submit":
        return "Auto Submitted"
    if status_key in {"submitted", "retaken"}:
        return "Self Submitted"
    if status_key == "retake_allowed":
        return "Retake Allowed"
    if status_key:
        return status_key.replace("_", " ").title()
    return "-"


def _topic_analysis_for_exam_result(result, question_meta):
    breakdown = getattr(result, "question_wise_breakdown", None)
    if not isinstance(breakdown, list) or not breakdown:
        return {"text": "-", "topics": []}

    by_topic = {}
    for item in breakdown:
        if not isinstance(item, dict):
            continue

        qid = str(item.get("question_id") or "")
        qmeta = question_meta.get(qid, {"topic": "General", "marks": 0.0})
        topic = (qmeta.get("topic") or "General").strip() or "General"

        if topic not in by_topic:
            by_topic[topic] = {
                "attempted": 0,
                "correct": 0,
                "marks_obtained": 0.0,
                "total_marks": 0.0,
                "total_questions": 0,
            }

        bucket = by_topic[topic]
        bucket["total_questions"] += 1
        bucket["total_marks"] += float(qmeta.get("marks") or 0)

        student_answer = item.get("student_answer")
        if student_answer not in (None, ""):
            bucket["attempted"] += 1
        if bool(item.get("correct")):
            bucket["correct"] += 1
        bucket["marks_obtained"] += float(item.get("marks_awarded") or 0)

    if not by_topic:
        return {"text": "-", "topics": []}

    def _pretty_num(value):
        try:
            num = float(value)
        except (TypeError, ValueError):
            return "0"
        if num.is_integer():
            return str(int(num))
        return f"{num:.2f}".rstrip("0").rstrip(".")

    chunks = []
    topics = []
    for topic in sorted(by_topic.keys()):
        d = by_topic[topic]
        unattempted = max(d["total_questions"] - d["attempted"], 0)
        marks_obtained = _pretty_num(d["marks_obtained"])
        total_marks = _pretty_num(d["total_marks"])
        topics.append(topic.title())
        chunks.append(
            f"{topic.title()} -> Attempted: {d['attempted']}, Unattempted: {unattempted}, Correct: {d['correct']}, Score: {marks_obtained}/{total_marks}"
        )
    return {"text": " ; ".join(chunks), "topics": topics}


def is_admin(user):
    return user.is_authenticated and user.role == 'admin'


def _build_admin_dashboard_stats():
    today = datetime.today().date()
    return {
        "students": StudentProfile.objects.count(),
        "colleges": College.objects.count(),
        "active_colleges": College.objects.filter(is_active=True).count(),
        "courses": Course.objects.count(),
        "sections": Section.objects.count(),
        "trainers": TrainerProfile.objects.count(),
        "domains": Domain.objects.count(),
        "subdomains": SubDomain.objects.count(),
        "exams": Exam.objects.count(),
        "scheduled_exams": ScheduledExam.objects.count(),
        "sessions": SessionSchedule.objects.count(),
        "today_sessions": SessionSchedule.objects.filter(date=today).count(),
        "feedback_entries": StudentFeedback.objects.count(),
    }

@login_required
@user_passes_test(is_admin)
def admin_dashboard_page(request):
    stats = _build_admin_dashboard_stats()

    return render(request, 'admin_dashboard.html', {
        'stats': stats,
    })


@login_required
@user_passes_test(is_admin)
@require_GET
def admin_dashboard_stats_api(request):
    return JsonResponse({
        "stats": _build_admin_dashboard_stats(),
        "updated_at": datetime.now().strftime("%H:%M:%S"),
    })

@login_required
@user_passes_test(is_admin)
def registration_page(request):
    return render(request, 'admin_registration.html')

@login_required
@user_passes_test(is_admin)
def admin_trainer_view_page(request):
    return render(request, 'admin_trainer.html')

@login_required
@user_passes_test(is_admin)
def admin_assessment_page(request):
    return render(request, 'admin_assessment.html')

@login_required
@user_passes_test(is_admin)
def admin_question_bank_page(request):
    return render(request, 'admin_practice_test.html')

@login_required
@user_passes_test(is_admin)
def admin_material_page(request):
    return render(request,'admin_material.html')

@login_required
@user_passes_test(is_admin)
def admin_trainer_page(request):
    return render(request, "admin_trainer.html")

@login_required
@user_passes_test(is_admin)
def admin_student_performance_page(request):
    return render(request, "admin_student_performance.html")

@login_required
@user_passes_test(is_admin)
def admin_student_resume_page(request):
    return render(request, 'admin_student_resume.html')

@login_required
@user_passes_test(is_admin)
def admin_schedule_page(request):
    return render(request, 'admin_schedule.html')

@login_required
@user_passes_test(is_admin)
def admin_attendance_page(request):
    return render(request, 'admin_attendance.html')

@login_required
@user_passes_test(is_admin)
def admin_feedback_page(request):
    return render(request, "admin_feedback.html")

@login_required
@user_passes_test(is_admin)
def admin_trainer_reports_page(request):
    return render(request, "admin_trainer_reports.html")


@login_required
@user_passes_test(is_admin)
def college_html_page(request):
    # change ordering to alphabetical for nicer UX; if you prefer newest first use '-created_at'
    colleges = College.objects.order_by('name')
    courses = Course.objects.order_by('name')
    sections = Section.objects.order_by('name')
    return render(request, 'admin_college.html', {
        'colleges': colleges,
        'courses': courses,
        'sections': sections,
    })

@login_required
@user_passes_test(is_admin)
def domain_module_page(request):
    
    return render (request,'domain_module.html' , {
    'domains': Domain.objects.all(),
    'subdomains': SubDomain.objects.select_related('domain'),
    'modules': Module.objects.select_related('domain', 'subdomain')
})

@login_required
@user_passes_test(is_admin)
def admin_registration(request):
    student_form = StudentRegistrationForm()
    bulk_form    = BulkStudentUploadForm()
    trainer_form = TrainerRegistrationForm()
    tpo_form     = TpoRegistrationForm()

    if request.method == 'POST':
        if 'student_single' in request.POST:
            student_form = StudentRegistrationForm(request.POST)
            if student_form.is_valid():
                student_form.save(request)
                messages.success(request, "Student created; verification email sent.")
                return redirect('admin_registration')

        elif 'student_bulk' in request.POST:
            bulk_form = BulkStudentUploadForm(request.POST, request.FILES)
            if bulk_form.is_valid():
                created, failed = bulk_form.process(request)

                # Success toast
                if created:
                    messages.success(
                        request,
                        f"Processed {len(created)} rows; verification emails sent to {', '.join(created)}."
                    )

                # Error toasts
                for err in failed:
                    messages.error(
                        request,
                        f"Row {err['row']} ({err['email']}): {err['error']}"
                    )

                return redirect('admin_registration')

        elif 'trainer' in request.POST:
            trainer_form = TrainerRegistrationForm(request.POST, request.FILES)
            if trainer_form.is_valid():
                trainer_form.save(request)
                messages.success(request, "Trainer created; verification email sent.")
                return redirect('admin_registration')

        elif 'tpo' in request.POST:
            tpo_form = TpoRegistrationForm(request.POST)
            if tpo_form.is_valid():
                tpo_form.save(request)
                messages.success(request, "TPO created; verification email sent.")
                return redirect('admin_registration')

    return render(request, 'admin_registration.html', {
        'student_form': student_form,
        'bulk_form':    bulk_form,
        'trainer_form': trainer_form,
        'tpo_form':     tpo_form,
    })


#---------- STUDENT_PERFORMANCE---------------#


@login_required
@user_passes_test(is_admin)
def admin_perf_dropdowns(request):
    """
    Returns colleges, courses, and year options.
    Optionally filter courses by ?college=<uuid> if the Course model supports it.
    """
    college_qs = College.objects.all().order_by("name")
    colleges = [{"id": str(c.id), "name": c.name} for c in college_qs]

    course_qs = Course.objects.all().order_by("name")
    college_id = request.GET.get("college")

    if college_id:
        course_qs = course_qs.filter(college_courses__college_id=college_id)

    # id + name are enough for the dropdown; don't reference non-existent fields
    courses = [{"id": str(c.id), "name": c.name} for c in course_qs]

    years = [{"value": i, "label": str(i)} for i in range(1, 6)]
    return JsonResponse({"colleges": colleges, "courses": courses, "years": years})

@login_required
@user_passes_test(is_admin)
def admin_perf_tests(request):
    """
    Return tests (PracticeTest/Exam) scheduled for the given filters:
    ?type=practice|exam&college=<id>&course=<id>&year=<int|list>&semester=<int|list>
    Shows only the exams/tests scheduled for that cohort.
    """
    ttype = request.GET.get("type")
    college_id = request.GET.get("college")
    course_raw = request.GET.get("course")
    year_raw = request.GET.get("year", "")
    semester_raw = request.GET.get("semester", "")

    # Validate required filters (Pre-assessment is open-guest and not scheduled to strict student cohorts)
    if ttype != "pre_assessment":
        if not (ttype and college_id and course_raw and year_raw and semester_raw):
            return JsonResponse({"tests": []})

        course_ids = [c.strip() for c in course_raw.split(",") if c.strip()]
        year_ids = [int(y.strip()) for y in year_raw.split(",") if y.strip().isdigit()]
        semester_ids = [int(s.strip()) for s in semester_raw.split(",") if s.strip().isdigit()]

        if not (course_ids and year_ids and semester_ids):
            return JsonResponse({"tests": []})

    tests = []

    # ---------------- PRACTICE TESTS ---------------- #
    if ttype == "practice":
        sched = (
            ScheduledPracticeTest.objects.filter(
                college_id=college_id,
                course_id__in=course_ids,
                year__in=year_ids,
                semester__in=semester_ids,
            )
            .select_related("practice_test")
            .order_by("practice_test__title")
        )

        seen = set()
        for s in sched:
            if s.practice_test_id not in seen:
                seen.add(s.practice_test_id)
                tests.append({
                    "id": str(s.practice_test_id),
                    "title": s.practice_test.title
                })

    # ---------------- EXAMS ---------------- #
    elif ttype == "exam":
        sched = (
            ScheduledExam.objects.filter(
                college_id=college_id,
                course_id__in=course_ids,
                year__in=year_ids,
                semester__in=semester_ids,
            )
            .select_related("exam")
            .order_by("exam__title")
        )

        seen = set()
        for s in sched:
            if s.exam_id not in seen:
                seen.add(s.exam_id)
                tests.append({
                    "id": str(s.exam_id),
                    "title": s.exam.title
                })

    # ---------------- ONE TIME ASSESSMENT / PRE-ASSESSMENT ---------------- #
    elif ttype == "pre_assessment":
        pre_exams = PreAssessmentExam.objects.all().select_related("exam").order_by("exam__title")
        for pe in pre_exams:
            tests.append({
                "id": str(pe.id),
                "title": f"{pe.exam.title} ({pe.code})"
            })

    return JsonResponse({"tests": tests})



@login_required
@user_passes_test(is_admin)
def admin_perf_results(request):
    """
    Return per-student results for Practice or Exam.
    Params:
      ?type=practice|exam&college=<id>&course=<id>&year=<int|list>&semester=<int|list>&test=<uuid|optional>
    """
    ttype = request.GET.get("type")
    college_id = request.GET.get("college")
    course_raw = request.GET.get("course", "")
    year_raw = request.GET.get("year", "")
    semester_raw = request.GET.get("semester", "")
    test_id = request.GET.get("test")  # optional

    if not (ttype and college_id and course_raw and year_raw):
        return JsonResponse({"results": [], "summary": {}, "chart": {} })

    course_ids = [c.strip() for c in course_raw.split(",") if c.strip()]
    year_ids = [int(y.strip()) for y in year_raw.split(",") if y.strip().isdigit()]
    semester_ids = [int(s.strip()) for s in semester_raw.split(",") if s.strip().isdigit()]

    if not (course_ids and year_ids):
        return JsonResponse({"results": [], "summary": {}, "chart": {} })

    section_raw = request.GET.get("section", "")
    section_names = [s.strip() for s in section_raw.split(",") if s.strip()]

    results_payload = []
    pass_count = 0
    fail_count = 0
    marks_list = []
    domain_scores = defaultdict(list)

    if ttype == "practice":
        # Limit to tests scheduled for that cohort
        sched = ScheduledPracticeTest.objects.filter(
            college_id=college_id, course_id__in=course_ids, year__in=year_ids
        )
        if semester_ids:
            sched = sched.filter(semester__in=semester_ids)
        sched = sched.values_list("practice_test_id", flat=True)

        if test_id:
            sched = sched.filter(id=test_id) if isinstance(sched, list) else [test_id]

        test_ids = list(set(sched))
        if not test_ids:
            return JsonResponse({"results": [], "summary": {}, "chart": {} })

        # Fetch practice results
        qs = (
            PracticeResult.objects
            .filter(
                practice_test_id__in=test_ids,
                student__college_id=college_id,
                student__course_id__in=course_ids,
                student__year__in=year_ids,
            )
        )
        if semester_ids:
            qs = qs.filter(student__semester__in=semester_ids)
        if section_names:
            qs = qs.filter(student__section__name__in=section_names)
            
        qs = (
            qs
            .select_related("student", "practice_test")
            .order_by("-submitted_at")
        )

        # Deduplicate to keep only the most recent attempt per student for this practice test
        grouped = defaultdict(list)
        for r in qs:
            grouped[(r.student_id, r.practice_test_id)].append(r)
        deduped = []
        for key, attempts in grouped.items():
            deduped.append(attempts[0])
        qs = deduped


        # Prepare map for passing marks + total per test
        tmap = {t.id: t for t in PracticeTest.objects.filter(id__in=test_ids)}
        practice_question_count = {
            q["practice_test_id"]: q["cnt"]
            for q in (
                PracticeQuestion.objects
                .filter(practice_test_id__in=test_ids)
                .values("practice_test_id")
                .annotate(cnt=Count("id"))
            )
        }

        for r in qs:
            test = tmap.get(r.practice_test_id)
            if not test:
                continue
            percentage = (r.marks_obtained / r.total_marks * 100.0) if r.total_marks else 0.0
            status = "Pass" if r.marks_obtained >= float(test.passing_marks) else "Fail"
            if status == "Pass":
                pass_count += 1
            else:
                fail_count += 1
            marks_list.append(percentage)
            if test and test.domain:
                dname = getattr(test.domain, "domain_name", "") or "General"
                domain_scores[dname].append(percentage)

            total_questions = int(practice_question_count.get(r.practice_test_id, 0) or 0)
            attempted = int(r.attempted_questions or 0)
            unattempted = max(total_questions - attempted, 0)

            results_payload.append({
                "result_id": r.id,
                "student_id": r.student_id,
                "usn": getattr(r.student, "usn", "") or "N/A",
                "student_name": _student_name(r.student),
                "college": getattr(getattr(r.student, "college", None), "name", ""),
                "course": getattr(getattr(r.student, "course", None), "name", ""),
                "title": test.title,
                "marks_obtained": r.marks_obtained,
                "total_marks": r.total_marks,
                "percentage": round(percentage, 2),
                "status": status,
                "attempted": attempted,
                "correct": r.correct_answers,
                "wrong": r.wrong_answers,
                "unattempted": unattempted,
                "tab_switch_count": "-",
                "time_spent_each_question": "-",
                "submission_mode": "Self Submitted",
                "topic_analysis": "-",
                "topic_list": ["General"],
                "time_taken": str(r.time_taken) if r.time_taken else "",
                "submitted_at": localtime(r.submitted_at).strftime("%Y-%m-%d %H:%M"),
            })

    elif ttype == "exam":
        sched = ScheduledExam.objects.filter(
            college_id=college_id, course_id__in=course_ids, year__in=year_ids
        )
        if semester_ids:
            sched = sched.filter(semester__in=semester_ids)
        sched = sched.values_list("exam_id", flat=True)

        if test_id:
            sched = sched.filter(id=test_id) if isinstance(sched, list) else [test_id]

        test_ids = list(set(sched))
        if not test_ids:
            return JsonResponse({"results": [], "summary": {}, "chart": {} })

        qs = (
            ExamResult.objects
            .filter(
                exam_id__in=test_ids,
                student__college_id=college_id,
                student__course_id__in=course_ids,
                student__year__in=year_ids,
            )
        )
        if semester_ids:
            qs = qs.filter(student__semester__in=semester_ids)
        if section_names:
            qs = qs.filter(student__section__name__in=section_names)
            
        qs = (
            qs
            .select_related("student", "exam")
            .order_by("-submitted_at")
        )

        # Deduplicate: prefer final/submitted attempts over active/in-progress, and keep only one per student/exam
        grouped = defaultdict(list)
        for r in qs:
            grouped[(r.student_id, r.exam_id)].append(r)
        deduped = []
        for key, attempts in grouped.items():
            final_attempts = [a for a in attempts if a.submission_status in ['submitted', 'accidental_submit', 'suspicious', 'retaken', 'retake_allowed']]
            if final_attempts:
                deduped.append(final_attempts[0])
            else:
                deduped.append(attempts[0])
        qs = deduped


        tmap = {t.id: t for t in Exam.objects.filter(id__in=test_ids)}
        question_meta_by_exam = defaultdict(dict)
        for q in ExamQuestion.objects.filter(exam_id__in=test_ids).values("id", "exam_id", "section_tag", "marks"):
            question_meta_by_exam[q["exam_id"]][str(q["id"])] = {
                "topic": (q.get("section_tag") or "General").strip() or "General",
                "marks": float(q.get("marks") or 0),
            }

        for r in qs:
            test = tmap.get(r.exam_id)
            if not test:
                continue
            percentage = (r.marks_obtained / r.total_marks * 100.0) if r.total_marks else 0.0
            status = "Pass" if r.marks_obtained >= float(test.passing_marks) else "Fail"
            if status == "Pass":
                pass_count += 1
            else:
                fail_count += 1
            marks_list.append(percentage)
            if test and test.domain:
                dname = getattr(test.domain, "domain_name", "") or "General"
                domain_scores[dname].append(percentage)

            attempted = int(r.attempted_questions or 0)
            total_questions = len(question_meta_by_exam.get(r.exam_id, {}))
            unattempted = max(total_questions - attempted, 0)
            tab_switch_count = int(getattr(r, "tab_switch_count", 0) or 0)
            submission_mode = _submission_mode_for_exam(r)
            time_spent_each_question = _format_question_time_map(getattr(r, "question_time_map", {}) or {})
            topic_summary = _topic_analysis_for_exam_result(r, question_meta_by_exam.get(r.exam_id, {}))
            topic_analysis = topic_summary["text"]
            topic_list = topic_summary["topics"]

            results_payload.append({
                "result_id": r.id,
                "student_id": r.student_id,
                "usn": getattr(r.student, "usn", "") or "N/A",
                "student_name": _student_name(r.student),
                "college": getattr(getattr(r.student, "college", None), "name", ""),
                "course": getattr(getattr(r.student, "course", None), "name", ""),
                "title": test.title,
                "marks_obtained": r.marks_obtained,
                "total_marks": r.total_marks,
                "percentage": round(percentage, 2),
                "status": status,
                "attempted": attempted,
                "correct": r.correct_answers,
                "wrong": r.wrong_answers,
                "unattempted": unattempted,
                "tab_switch_count": tab_switch_count,
                "time_spent_each_question": time_spent_each_question,
                "submission_mode": submission_mode,
                "topic_analysis": topic_analysis,
                "topic_list": topic_list,
                "time_taken": str(r.time_taken) if r.time_taken else "",
                "submitted_at": localtime(r.submitted_at).strftime("%Y-%m-%d %H:%M"),
                "device_info": getattr(r, "device_info", "") or "-",
                "device_type": getattr(r, "device_type", "") or "-",
                "os_name": getattr(r, "os_name", "") or "-",
                "browser_name": getattr(r, "browser_name", "") or "-",
                "ip_address": getattr(r, "ip_address", "") or "-",
            })

    elif ttype == "pre_assessment":
        # 1. Translate college and course UUID/IDs to text values for guest match
        college_name = College.objects.filter(id=college_id).values_list("name", flat=True).first() or ""
        course_names = list(Course.objects.filter(id__in=course_ids).values_list("name", flat=True))

        # Auto-evaluate and submit any pending in-progress attempts
        unsubmitted = ExamResult.objects.filter(
            pre_assessment__isnull=False
        ).exclude(submission_status__in=['submitted', 'accidental_submit', 'retaken'])
        for un_res in unsubmitted:
            try:
                from pre_assessment.views import evaluate_and_finalize_pre_assessment_result
                evaluate_and_finalize_pre_assessment_result(un_res)
            except Exception:
                pass

        qs = ExamResult.objects.filter(pre_assessment__isnull=False)

        if test_id:
            qs = qs.filter(pre_assessment_id=test_id)

        if college_name:
            qs = qs.filter(candidate_college=college_name)
        if course_names:
            qs = qs.filter(candidate_course__in=course_names)
        if year_ids:
            qs = qs.filter(candidate_year__in=year_ids)
        if semester_ids:
            qs = qs.filter(candidate_sem__in=semester_ids)

        # Fetch and deduplicate to get the latest attempt per candidate email
        qs = qs.select_related("pre_assessment", "pre_assessment__exam").order_by("-submitted_at")
        grouped = defaultdict(list)
        for r in qs:
            grouped[(r.candidate_email, r.pre_assessment_id)].append(r)
        deduped = []
        for key, attempts in grouped.items():
            final_attempts = [a for a in attempts if a.submission_status in ['submitted', 'accidental_submit', 'suspicious', 'retaken', 'retake_allowed']]
            if final_attempts:
                deduped.append(final_attempts[0])
            else:
                deduped.append(attempts[0])
        qs = deduped

        # Map exams metadata
        pre_exam_ids = list(set([r.pre_assessment_id for r in qs]))
        pre_exams = {pe.id: pe for pe in PreAssessmentExam.objects.filter(id__in=pre_exam_ids).select_related("exam")}
        
        exam_ids = list(set([pe.exam_id for pe in pre_exams.values()]))
        tmap = {pe.id: pe.exam for pe in pre_exams.values()}
        
        question_meta_by_exam = defaultdict(dict)
        for q in ExamQuestion.objects.filter(exam_id__in=exam_ids).values("id", "exam_id", "section_tag", "marks"):
            question_meta_by_exam[q["exam_id"]][str(q["id"])] = {
                "topic": (q.get("section_tag") or "General").strip() or "General",
                "marks": float(q.get("marks") or 0),
            }

        # Populate payload
        for r in qs:
            pe_obj = pre_exams.get(r.pre_assessment_id)
            if not pe_obj:
                continue
            test = pe_obj.exam
            
            percentage = (r.marks_obtained / r.total_marks * 100.0) if r.total_marks else 0.0
            status = "Pass" if r.marks_obtained >= float(test.passing_marks) else "Fail"
            if status == "Pass":
                pass_count += 1
            else:
                fail_count += 1
            marks_list.append(percentage)
            if test and test.domain:
                dname = getattr(test.domain, "domain_name", "") or "General"
                domain_scores[dname].append(percentage)

            attempted = int(r.attempted_questions or 0)
            total_questions = len(question_meta_by_exam.get(test.id, {}))
            unattempted = max(total_questions - attempted, 0)
            tab_switch_count = int(getattr(r, "tab_switch_count", 0) or 0)
            submission_mode = _submission_mode_for_exam(r)
            time_spent_each_question = _format_question_time_map(getattr(r, "question_time_map", {}) or {})
            topic_summary = _topic_analysis_for_exam_result(r, question_meta_by_exam.get(test.id, {}))
            topic_analysis = topic_summary["text"]
            topic_list = topic_summary["topics"]

            results_payload.append({
                "result_id": r.id,
                "student_id": None,
                "usn": getattr(r, "candidate_usn", "") or "N/A",
                "student_name": r.candidate_name or "Guest",
                "email": getattr(r, "candidate_email", "") or "-",
                "phone": getattr(r, "candidate_phone", "") or "-",
                "college": r.candidate_college or "-",
                "course": r.candidate_course or "-",
                "title": f"{test.title} ({pe_obj.code})",
                "custom_data": getattr(r, "candidate_custom_data", {}) or {},
                "marks_obtained": r.marks_obtained,
                "total_marks": r.total_marks,
                "percentage": round(percentage, 2),
                "status": status,
                "attempted": attempted,
                "correct": r.correct_answers,
                "wrong": r.wrong_answers,
                "unattempted": unattempted,
                "tab_switch_count": tab_switch_count,
                "time_spent_each_question": time_spent_each_question,
                "submission_mode": submission_mode,
                "topic_analysis": topic_analysis,
                "topic_list": topic_list,
                "time_taken": str(r.time_taken) if r.time_taken else "",
                "submitted_at": localtime(r.submitted_at).strftime("%Y-%m-%d %H:%M"),
                "device_info": getattr(r, "device_info", "") or "-",
                "device_type": getattr(r, "device_type", "") or "-",
                "os_name": getattr(r, "os_name", "") or "-",
                "browser_name": getattr(r, "browser_name", "") or "-",
                "ip_address": getattr(r, "ip_address", "") or "-",
            })

        test_ids = list(set([pe.exam_id for pe in pre_exams.values()]))

    # Summary & chart
    avg = round(sum(marks_list) / len(marks_list), 2) if marks_list else 0.0
    best = round(max(marks_list), 2) if marks_list else 0.0
    worst = round(min(marks_list), 2) if marks_list else 0.0

    topic_to_students = defaultdict(set)
    for row in results_payload:
        student_name = row.get("student_name", "")
        for topic in row.get("topic_list", []):
            topic_name = str(topic or "").strip()
            if topic_name:
                topic_to_students[topic_name].add(student_name)

    topic_items = sorted(
        topic_to_students.items(),
        key=lambda item: (-len(item[1]), item[0].lower())
    )

    # Chart.js payload
    chart = {
        "labels": [f"{r['student_name']}" for r in results_payload],
        "values": [r["percentage"] for r in results_payload],
        "topic_labels": [topic for topic, _students in topic_items],
        "topic_values": [len(students) for _topic, students in topic_items],
        "topic_students": {
            topic: sorted(list(students))
            for topic, students in topic_items
        },
    }

    # Aggregate Coding Metrics for Coding Analytics Dashboard
    coding_attempts = 0
    coding_passed_cases = 0
    coding_total_cases = 0
    languages_map = defaultdict(int)
    verdicts_map = defaultdict(int)

    for r in qs:
        breakdown = r.question_wise_breakdown or []
        if not isinstance(breakdown, list):
            continue
        for q_attempt in breakdown:
            ans_str = q_attempt.get("student_answer", "")
            is_code = False
            lang = None
            if ans_str:
                try:
                    import json
                    ans_data = json.loads(ans_str)
                    if isinstance(ans_data, dict) and "language" in ans_data:
                        is_code = True
                        lang = ans_data.get("language")
                except Exception:
                    pass
            if not is_code and ("passed_cases" in q_attempt or "verdict" in q_attempt):
                is_code = True

            if is_code:
                coding_attempts += 1
                if lang:
                    languages_map[str(lang).lower()] += 1
                
                verdict = q_attempt.get("verdict")
                if verdict:
                    verdicts_map[str(verdict)] += 1
                
                passed = q_attempt.get("passed_cases", 0)
                total = q_attempt.get("total_cases", 0)
                try:
                    coding_passed_cases += int(passed or 0)
                    coding_total_cases += int(total or 0)
                except Exception:
                    pass

    # Calculate test summary stats
    test_negative_marking = "No"
    total_questions_count = 0
    test_duration = 0
    test_max_marks = 0
    
    if test_ids:
        test_obj = next(iter(tmap.values())) if tmap else None
        if test_obj:
            test_duration = getattr(test_obj, "duration_minutes", 0) or 0
            test_max_marks = getattr(test_obj, "max_marks", 0) or 0
            
        if ttype == "practice":
            total_questions_count = sum(practice_question_count.values())
            has_negative = PracticeQuestion.objects.filter(practice_test_id__in=test_ids, negative_mark__gt=0).exists()
            test_negative_marking = "Yes" if has_negative else "No"
        else:
            total_questions_count = len(question_meta_by_exam.get(test_ids[0], {})) if test_ids else 0
            has_negative = ExamQuestion.objects.filter(exam_id__in=test_ids, negative_mark__gt=0).exists()
            test_negative_marking = "Yes" if has_negative else "No"

    # Calculate average time spent
    time_secs = []
    for r in qs:
        t_taken = getattr(r, "time_taken", None)
        if t_taken:
            if hasattr(t_taken, "total_seconds"):
                time_secs.append(t_taken.total_seconds())
    
    avg_secs = sum(time_secs) / len(time_secs) if time_secs else 0.0
    avg_minutes = int(avg_secs // 60)
    avg_remaining_secs = int(avg_secs % 60)
    avg_time_taken_formatted = f"{avg_minutes:02d}:{avg_remaining_secs:02d}"

    summary = {
        "pass": pass_count,
        "fail": fail_count,
        "avg": avg,
        "best": best,
        "worst": worst,
        "test_metadata": {
            "total_questions": total_questions_count,
            "max_marks": test_max_marks,
            "duration": test_duration,
            "negative_marking": test_negative_marking
        },
        "avg_time_taken_formatted": avg_time_taken_formatted,
        "coding": {
            "attempts": coding_attempts,
            "passed_cases": coding_passed_cases,
            "total_cases": coding_total_cases,
            "languages": dict(languages_map),
            "verdicts": dict(verdicts_map)
        }
    }
    domains_payload = [
        {"domain": name, "avg": round(sum(scores) / len(scores), 2)}
        for name, scores in domain_scores.items()
    ]
    return JsonResponse({"results": results_payload, "summary": summary, "chart": chart, "domains": domains_payload})


def is_admin_or_tpo(user):
    return user.is_authenticated and user.role in ('admin', 'tpo')


@login_required
@user_passes_test(is_admin_or_tpo)
def admin_perf_result_detail(request):
    """
    Returns detailed question-by-question result for an ExamResult or PracticeResult.
    """
    ttype = request.GET.get("type")
    result_id = request.GET.get("id")

    if not ttype or not result_id:
        return JsonResponse({"error": "Missing parameters"}, status=400)

    try:
        if ttype in ("exam", "pre_assessment"):
            result = get_object_or_404(ExamResult.objects.select_related("student", "exam", "pre_assessment", "pre_assessment__exam"), id=result_id)
            exam_obj = result.exam or (result.pre_assessment.exam if result.pre_assessment else None)
            test_title = result.exam_title or (exam_obj.title if exam_obj else "N/A")
            passing_marks = float(exam_obj.passing_marks) if (exam_obj and exam_obj.passing_marks) else 0.0
            
            # Fetch exam questions to map the text
            exam_questions = list(ExamQuestion.objects.filter(exam=exam_obj)) if exam_obj else []
            question_map = {str(q.id): q for q in exam_questions}
            
            breakdown_map = {}
            for item in (result.question_wise_breakdown or []):
                if isinstance(item, dict):
                    q_id = str(item.get("question_id") or "")
                    if q_id:
                        breakdown_map[q_id] = item

            breakdown_data = []
            seen_ids = set()

            # Iterate through exam_questions in natural exam order
            for q in exam_questions:
                q_id = str(q.id)
                seen_ids.add(q_id)
                item = breakdown_map.get(q_id)

                if item is not None:
                    ans_str = item.get("student_answer") or item.get("answer") or ""
                    parsed_code = None
                    if ans_str:
                        try:
                            import json
                            parsed = json.loads(ans_str)
                            if isinstance(parsed, dict) and "code" in parsed:
                                parsed_code = {
                                    "code": parsed.get("code", ""),
                                    "language": parsed.get("language", "python")
                                }
                        except Exception:
                            pass

                    breakdown_data.append({
                        "question_id": q_id,
                        "question_text": q.question_text if q else "Question text not available",
                        "type": item.get("type") or (q.type if q else "N/A"),
                        "student_answer": ans_str,
                        "parsed_code": parsed_code,
                        "correct_answer": item.get("correct_answer") or (q.correct_answer if q else "N/A"),
                        "correct": item.get("correct", False),
                        "marks_awarded": item.get("marks_awarded", 0),
                        "verdict": item.get("verdict") or ("Accepted" if item.get("correct") else "Wrong Answer"),
                        "passed": item.get("passed", 0),
                        "total": item.get("total", 0),
                        "test_results": item.get("test_results") or [],
                    })
                else:
                    # Fallback for historical results where only some questions were captured in breakdown JSON
                    breakdown_data.append({
                        "question_id": q_id,
                        "question_text": q.question_text,
                        "type": q.type,
                        "student_answer": "— (Recorded in final score)" if (result.attempted_questions or 0) > 0 else "— (Unanswered)",
                        "parsed_code": None,
                        "correct_answer": q.correct_answer if q.type in ["MCQ", "TF"] else "—",
                        "correct": False,
                        "marks_awarded": 0,
                        "verdict": "Completed",
                        "passed": 0,
                        "total": 0,
                        "test_results": [],
                    })

            # Append any leftover breakdown items (e.g. dynamic questions)
            for item in (result.question_wise_breakdown or []):
                if isinstance(item, dict):
                    q_id = str(item.get("question_id") or "")
                    if q_id and q_id not in seen_ids:
                        q_obj = question_map.get(q_id)
                        ans_str = item.get("student_answer") or item.get("answer") or ""
                        parsed_code = None
                        if ans_str:
                            try:
                                import json
                                parsed = json.loads(ans_str)
                                if isinstance(parsed, dict) and "code" in parsed:
                                    parsed_code = {
                                        "code": parsed.get("code", ""),
                                        "language": parsed.get("language", "python")
                                    }
                            except Exception:
                                pass

                        breakdown_data.append({
                            "question_id": q_id,
                            "question_text": q_obj.question_text if q_obj else "Question text not available",
                            "type": item.get("type") or (q_obj.type if q_obj else "N/A"),
                            "student_answer": ans_str,
                            "parsed_code": parsed_code,
                            "correct_answer": item.get("correct_answer") or (q_obj.correct_answer if q_obj else "N/A"),
                            "correct": item.get("correct", False),
                            "marks_awarded": item.get("marks_awarded", 0),
                            "verdict": item.get("verdict") or ("Accepted" if item.get("correct") else "Wrong Answer"),
                            "passed": item.get("passed", 0),
                            "total": item.get("total", 0),
                            "test_results": item.get("test_results") or [],
                        })

        elif ttype == "practice":
            result = get_object_or_404(PracticeResult.objects.select_related("student", "practice_test"), id=result_id)
            test_title = result.practice_title or (result.practice_test.title if result.practice_test else "N/A")
            passing_marks = float(result.practice_test.passing_marks) if (result.practice_test and result.practice_test.passing_marks) else 0.0
            
            # Fetch practice questions
            practice_questions = list(PracticeQuestion.objects.filter(practice_test=result.practice_test)) if result.practice_test else []
            question_map = {str(q.id): q for q in practice_questions}
            
            breakdown_map = {}
            for item in (result.question_wise_breakdown or []):
                if isinstance(item, dict):
                    q_id = str(item.get("question_id") or "")
                    if q_id:
                        breakdown_map[q_id] = item

            breakdown_data = []
            seen_ids = set()

            for q in practice_questions:
                q_id = str(q.id)
                seen_ids.add(q_id)
                item = breakdown_map.get(q_id)

                if item is not None:
                    ans_str = item.get("student_answer") or item.get("answer") or ""
                    parsed_code = None
                    if ans_str:
                        try:
                            import json
                            parsed = json.loads(ans_str)
                            if isinstance(parsed, dict) and "code" in parsed:
                                parsed_code = {
                                    "code": parsed.get("code", ""),
                                    "language": parsed.get("language", "python")
                                }
                        except Exception:
                            pass

                    breakdown_data.append({
                        "question_id": q_id,
                        "question_text": q.question_text if q else "Question text not available",
                        "type": item.get("type") or (q.type if q else "N/A"),
                        "student_answer": ans_str,
                        "parsed_code": parsed_code,
                        "correct_answer": item.get("correct_answer") or (q.correct_answer if q else "N/A"),
                        "correct": item.get("correct", False),
                        "marks_awarded": item.get("marks_awarded", 0),
                        "verdict": item.get("verdict") or ("Accepted" if item.get("correct") else "Wrong Answer"),
                        "passed": item.get("passed", 0),
                        "total": item.get("total", 0),
                    })
                else:
                    breakdown_data.append({
                        "question_id": q_id,
                        "question_text": q.question_text,
                        "type": q.type,
                        "student_answer": "— (Unanswered)",
                        "parsed_code": None,
                        "correct_answer": q.correct_answer if q.type in ["MCQ", "TF"] else "—",
                        "correct": False,
                        "marks_awarded": 0,
                        "verdict": "Completed",
                        "passed": 0,
                        "total": 0,
                    })

            for item in (result.question_wise_breakdown or []):
                if isinstance(item, dict):
                    q_id = str(item.get("question_id") or "")
                    if q_id and q_id not in seen_ids:
                        q_obj = question_map.get(q_id)
                        ans_str = item.get("student_answer") or item.get("answer") or ""
                        parsed_code = None
                        if ans_str:
                            try:
                                import json
                                parsed = json.loads(ans_str)
                                if isinstance(parsed, dict) and "code" in parsed:
                                    parsed_code = {
                                        "code": parsed.get("code", ""),
                                        "language": parsed.get("language", "python")
                                    }
                            except Exception:
                                pass

                        breakdown_data.append({
                            "question_id": q_id,
                            "question_text": q_obj.question_text if q_obj else "Question text not available",
                            "type": item.get("type") or (q_obj.type if q_obj else "N/A"),
                            "student_answer": ans_str,
                            "parsed_code": parsed_code,
                            "correct_answer": item.get("correct_answer") or (q_obj.correct_answer if q_obj else "N/A"),
                            "correct": item.get("correct", False),
                            "marks_awarded": item.get("marks_awarded", 0),
                            "verdict": item.get("verdict") or ("Accepted" if item.get("correct") else "Wrong Answer"),
                            "passed": item.get("passed", 0),
                            "total": item.get("total", 0),
                        })
        else:
            return JsonResponse({"error": "Invalid test type"}, status=400)

        # Basic student details
        student = result.student
        if student:
            student_data = {
                "name": _student_name(student),
                "usn": student.usn or "N/A",
                "email": student.user.email if (student.user and student.user.email) else "N/A",
                "phone": getattr(student, "phone", "N/A") or "N/A",
                "college": student.college.name if student.college else "N/A",
                "course": student.course.name if student.course else "N/A",
                "semester": student.semester,
                "year": student.year,
                "section": student.section.name if student.section else "N/A",
                "custom_data": {},
            }
        else:
            student_data = {
                "name": getattr(result, "candidate_name", "N/A") or "N/A",
                "usn": getattr(result, "candidate_usn", "N/A") or "N/A",
                "email": getattr(result, "candidate_email", "N/A") or "N/A",
                "phone": getattr(result, "candidate_phone", "N/A") or "N/A",
                "college": getattr(result, "candidate_college", "N/A") or "N/A",
                "course": getattr(result, "candidate_course", "N/A") or "N/A",
                "semester": getattr(result, "candidate_sem", "N/A") or "N/A",
                "year": getattr(result, "candidate_year", "N/A") or "N/A",
                "section": "N/A",
                "custom_data": getattr(result, "candidate_custom_data", {}) or {},
            }

        # Format time taken
        time_taken_str = "N/A"
        if result.time_taken:
            total_seconds = int(result.time_taken.total_seconds())
            mins = total_seconds // 60
            secs = total_seconds % 60
            time_taken_str = f"{mins}m {secs}s"

        response_data = {
            "student": student_data,
            "test_title": test_title,
            "test_type": ttype.upper(),
            "total_marks": result.total_marks,
            "marks_obtained": result.marks_obtained,
            "passing_marks": passing_marks,
            "percentage": round((result.marks_obtained / result.total_marks * 100.0) if result.total_marks else 0.0, 2),
            "status": "Pass" if result.marks_obtained >= passing_marks else "Fail",
            "attempted_questions": result.attempted_questions,
            "correct_answers": result.correct_answers,
            "wrong_answers": result.wrong_answers,
            "time_taken": time_taken_str,
            "submitted_at": localtime(result.submitted_at).strftime("%Y-%m-%d %I:%M %p"),
            "question_wise_breakdown": breakdown_data,
        }

        if ttype in ("exam", "pre_assessment"):
            response_data.update({
                "tab_switch_count": getattr(result, "tab_switch_count", 0),
                "submission_mode": _submission_mode_for_exam(result),
                "device_info": getattr(result, "device_info", "") or "-",
                "device_type": getattr(result, "device_type", "") or "-",
                "os_name": getattr(result, "os_name", "") or "-",
                "browser_name": getattr(result, "browser_name", "") or "-",
                "ip_address": getattr(result, "ip_address", "") or "-",
            })

        return JsonResponse(response_data)

    except Exception as e:
        import logging
        logger = logging.getLogger(__name__)
        logger.exception("Error in admin_perf_result_detail API")
        return JsonResponse({"error": str(e)}, status=500)


@login_required
@user_passes_test(is_admin)
def admin_perf_download(request):
    """
    Download Excel report for student performance with enriched monitoring columns and custom collected student data.
    """
    from django.http import HttpResponse
    from django.test.client import RequestFactory
    from io import BytesIO
    from datetime import datetime
    import json
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image
    from openpyxl.chart import BarChart, Reference, PieChart
    from openpyxl.chart.label import DataLabelList
    from openpyxl.utils import get_column_letter
    from student.models import StudentProfile

    # Validate type
    ttype = request.GET.get("type")
    if ttype not in ("practice", "exam", "pre_assessment"):
        raise Http404("Invalid type")

    # Fetch data from admin_perf_results
    rf = RequestFactory()
    fake = rf.get("/fake", data=request.GET)
    fake.user = request.user
    res = admin_perf_results(fake)
    data = json.loads(res.content.decode("utf-8"))

    results = data.get("results", [])
    summary = data.get("summary", {})

    # Workbook setup
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Performance"

    # ---------- HEADER ----------
    green_fill = PatternFill(start_color="008037", end_color="008037", fill_type="solid")

    # Logo in corner
    logo_path = "static/images/atomshaalelogo.png"
    try:
        img = Image(logo_path)
        img.width, img.height = 160, 70
        ws.add_image(img, "A1")
    except Exception:
        pass

    # Title (green banner only for middle section)
    ws.merge_cells("C1:O3")
    ws["C1"] = "STUDENT PERFORMANCE"
    ws["C1"].font = Font(size=22, bold=True, color="FFFFFF", name="Calibri")
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")
    for row in ws["C1:O3"]:
        for cell in row:
            cell.fill = green_fill

    # ---------- SUMMARY ----------
    start_row = 5
    ws[f"A{start_row}"] = "Summary"
    ws[f"A{start_row}"].font = Font(size=14, bold=True)
    summary_data = [
        ["Total Students", len(results)],
        ["Passed", summary.get("pass", 0)],
        ["Failed", summary.get("fail", 0)],
        ["Average (%)", summary.get("avg", 0)],
        ["Best (%)", summary.get("best", 0)],
        ["Lowest (%)", summary.get("worst", 0)],
    ]
    for label, value in summary_data:
        start_row += 1
        ws[f"A{start_row}"] = label
        ws[f"B{start_row}"] = value

    start_row += 2

    # Discover custom collected student data columns
    custom_keys = []
    has_email = any(bool(r.get("email") and r.get("email") != "-") for r in results)
    has_phone = any(bool(r.get("phone") and r.get("phone") != "-") for r in results)

    for r in results:
        cdata = r.get("custom_data") or {}
        if isinstance(cdata, dict):
            for k in cdata.keys():
                k_clean = str(k).strip()
                if k_clean and k_clean not in custom_keys:
                    custom_keys.append(k_clean)

    # ---------- TABLE HEADER ----------
    headers = [
        "USN", "Student Name", "College", "Course", "Test Title"
    ]
    if has_email:
        headers.append("Email")
    if has_phone:
        headers.append("Phone")
    for ck in custom_keys:
        headers.append(ck)

    headers.extend([
        "Marks Obtained", "Total Marks", "Percentage", "Status",
        "Attempted", "Correct", "Wrong", "Unattempted", "Tab Switch",
        "Time Spent (Each Question)", "Submission Mode", "Section/Topic Analysis",
        "Time Taken", "Submitted At", "Device / OS", "Browser", "IP Address"
    ])
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="008037", end_color="008037", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    for i, header in enumerate(headers, start=1):
        c = ws.cell(row=ws.max_row, column=i)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = thin_border

    # ---------- TABLE DATA ----------
    for r in results:
        usn = r.get("usn")
        if not usn or usn == "N/A":
            if r.get("student_id"):
                try:
                    student_profile = StudentProfile.objects.filter(id=r["student_id"]).first()
                    usn = student_profile.usn if student_profile else "N/A"
                except Exception:
                    usn = "N/A"
            elif r.get("result_id"):
                try:
                    exam_res = ExamResult.objects.filter(id=r["result_id"]).first()
                    if exam_res:
                        usn = exam_res.candidate_usn or (exam_res.student.usn if exam_res.student else "N/A")
                    else:
                        usn = "N/A"
                except Exception:
                    usn = "N/A"
            else:
                usn = "N/A"

        attempted = int(r.get("attempted", 0) or 0)
        unattempted = int(r.get("unattempted", 0) or 0)

        row = [
            usn,
            r["student_name"],
            r["college"],
            r["course"],
            r["title"],
        ]
        if has_email:
            row.append(r.get("email") or "-")
        if has_phone:
            row.append(r.get("phone") or "-")
        for ck in custom_keys:
            cdata = r.get("custom_data") or {}
            row.append(cdata.get(ck, "-") if isinstance(cdata, dict) else "-")

        row.extend([
            r["marks_obtained"],
            r["total_marks"],
            r["percentage"],
            r["status"],
            attempted,
            r["correct"],
            r["wrong"],
            unattempted,
            r.get("tab_switch_count", "-"),
            r.get("time_spent_each_question", "-"),
            r.get("submission_mode", "-"),
            r.get("topic_analysis", "-"),
            r["time_taken"],
            r["submitted_at"],
            r.get("device_info", "-"),
            r.get("browser_name", "-"),
            r.get("ip_address", "-"),
        ])
        ws.append(row)

    # ---------- AUTO-FIT COLUMNS ----------
    for i, column_cells in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_length + 3, 12), 40)

    # ---------- ALIGNMENT ----------
    for row in ws.iter_rows(min_row=ws.min_row + 12, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center", horizontal="center")
            cell.border = thin_border

    # ---------- FOOTER ----------
    footer_row = ws.max_row + 3
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=len(headers))
    footer = ws[f"A{footer_row}"]
    footer.value = "Atomshaale — Powered by ATOM"
    footer.font = Font(italic=True, color="808080", size=10)
    footer.alignment = Alignment(horizontal="center", vertical="center")

    # ---------- CHARTS SHEET ----------
    if results:
        chart_ws = wb.create_sheet("Charts")
        chart_ws.append(["Student", "Percentage"])
        for r in results:
            chart_ws.append([r["student_name"], r["percentage"]])

        # Bar Chart - Student vs Percentage
        bar = BarChart()
        bar.title = "Student Performance (%)"
        data = Reference(chart_ws, min_col=2, min_row=1, max_row=len(results) + 1)
        cats = Reference(chart_ws, min_col=1, min_row=2, max_row=len(results) + 1)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        bar.y_axis.title = "Percentage"
        bar.x_axis.title = "Students"
        bar.dataLabels = DataLabelList(showVal=True)
        chart_ws.add_chart(bar, "E2")

        # Pie Chart - Pass vs Fail
        chart_ws["H3"], chart_ws["H4"] = "Pass", "Fail"
        chart_ws["I3"], chart_ws["I4"] = summary.get("pass", 0), summary.get("fail", 0)
        pie = PieChart()
        pie.title = "Pass vs Fail"
        pie.add_data(Reference(chart_ws, min_col=9, min_row=3, max_row=4))
        pie.set_categories(Reference(chart_ws, min_col=8, min_row=3, max_row=4))
        pie.dataLabels = DataLabelList(showPercent=True)
        chart_ws.add_chart(pie, "E20")

    # ---------- EXPORT ----------
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"{ttype}_student_performance_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    response = HttpResponse(
        buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename=\"{filename}\"'
    return response

    #--------------student resume ---------------#

@login_required
@user_passes_test(is_admin)
def admin_student_resume_api(request):
    filters = {}
    colleges = request.GET.getlist('college') or request.GET.get('college')
    if colleges:
        college_ids = []
        if isinstance(colleges, list):
            for val in colleges:
                college_ids.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            college_ids = [v.strip() for v in colleges.split(',') if v.strip()]
        if college_ids:
            filters['college_id__in'] = college_ids

    courses = request.GET.getlist('course') or request.GET.get('course')
    if courses:
        course_ids = []
        if isinstance(courses, list):
            for val in courses:
                course_ids.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            course_ids = [v.strip() for v in courses.split(',') if v.strip()]
        if course_ids:
            filters['course_id__in'] = course_ids

    semesters = request.GET.getlist('semester') or request.GET.get('semester')
    if semesters:
        sem_list = []
        if isinstance(semesters, list):
            for val in semesters:
                sem_list.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            sem_list = [v.strip() for v in semesters.split(',') if v.strip()]
        if sem_list:
            filters['semester__in'] = sem_list

    years = request.GET.getlist('year') or request.GET.get('year')
    if years:
        year_list = []
        if isinstance(years, list):
            for val in years:
                year_list.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            year_list = [v.strip() for v in years.split(',') if v.strip()]
        if year_list:
            filters['year__in'] = year_list

    q_term = (request.GET.get('q') or '').strip()

    qs = (
        StudentProfile.objects
        .filter(**filters)
        .filter(resume__isnull=False)
        .exclude(resume='')
        .select_related('user', 'college', 'course')
    )

    if q_term:
        qs = qs.filter(
            Q(usn__icontains=q_term) |
            Q(user__first_name__icontains=q_term) |
            Q(user__last_name__icontains=q_term) |
            Q(user__username__icontains=q_term) |
            Q(user__email__icontains=q_term) |
            Q(college__name__icontains=q_term) |
            Q(course__name__icontains=q_term)
        )

    qs = qs.order_by('usn')

    data = []
    for s in qs:
        # Robust full name
        name = ''
        try:
            if hasattr(s.user, 'get_full_name'):
                name = (s.user.get_full_name() or '').strip()
            if not name:
                first = getattr(s.user, 'first_name', '') or ''
                last = getattr(s.user, 'last_name', '') or ''
                name = (f'{first} {last}').strip() or getattr(s.user, 'username', '')
        except Exception:
            name = getattr(s.user, 'username', '') or '—'

        # Safe URL
        resume_url = ''
        try:
            if s.resume and getattr(s.resume, 'url', None):
                resume_url = s.resume.url
        except Exception:
            resume_url = ''

        data.append({
            'usn': s.usn,
            'name': name,
            'email': getattr(s.user, 'email', '') or '',
            'resume_url': resume_url,
        })

    return JsonResponse(data, safe=False)


@login_required
@user_passes_test(is_admin)
def admin_college_list_api(request):
    colleges = College.objects.order_by('name').values('id', 'name')
    return JsonResponse(list(colleges), safe=False)


@login_required
@user_passes_test(is_admin)
def admin_course_list_api(request):
    courses = Course.objects.order_by('name').values('id', 'name')
    return JsonResponse(list(courses), safe=False)


@login_required
@user_passes_test(is_admin)
def admin_student_resume_filters_api(request):
    semesters = (
        StudentProfile.objects
        .filter(semester__isnull=False)
        .values_list('semester', flat=True)
        .distinct()
        .order_by('semester')
    )
    years = (
        StudentProfile.objects
        .filter(year__isnull=False)
        .values_list('year', flat=True)
        .distinct()
        .order_by('year')
    )
    return JsonResponse({
        'semesters': list(semesters),
        'years': list(years),
    })


@login_required
@user_passes_test(is_admin)
def admin_student_resume_download_zip(request):
    import zipfile
    import os
    from io import BytesIO
    from django.http import HttpResponse

    filters = {}
    colleges = request.GET.getlist('college') or request.GET.get('college')
    if colleges:
        college_ids = []
        if isinstance(colleges, list):
            for val in colleges:
                college_ids.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            college_ids = [v.strip() for v in colleges.split(',') if v.strip()]
        if college_ids:
            filters['college_id__in'] = college_ids

    courses = request.GET.getlist('course') or request.GET.get('course')
    if courses:
        course_ids = []
        if isinstance(courses, list):
            for val in courses:
                course_ids.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            course_ids = [v.strip() for v in courses.split(',') if v.strip()]
        if course_ids:
            filters['course_id__in'] = course_ids

    semesters = request.GET.getlist('semester') or request.GET.get('semester')
    if semesters:
        sem_list = []
        if isinstance(semesters, list):
            for val in semesters:
                sem_list.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            sem_list = [v.strip() for v in semesters.split(',') if v.strip()]
        if sem_list:
            filters['semester__in'] = sem_list

    years = request.GET.getlist('year') or request.GET.get('year')
    if years:
        year_list = []
        if isinstance(years, list):
            for val in years:
                year_list.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            year_list = [v.strip() for v in years.split(',') if v.strip()]
        if year_list:
            filters['year__in'] = year_list

    q_term = (request.GET.get('q') or '').strip()
    selected_usns = request.GET.get('usns') # comma-separated list of USNs

    qs = (
        StudentProfile.objects
        .filter(**filters)
        .filter(resume__isnull=False)
        .exclude(resume='')
        .select_related('user')
    )

    if selected_usns:
        usn_list = [u.strip() for u in selected_usns.split(',') if u.strip()]
        if usn_list:
            qs = qs.filter(usn__in=usn_list)

    if q_term and not selected_usns:
        qs = qs.filter(
            Q(usn__icontains=q_term) |
            Q(user__first_name__icontains=q_term) |
            Q(user__last_name__icontains=q_term) |
            Q(user__username__icontains=q_term) |
            Q(user__email__icontains=q_term)
        )

    if not qs.exists():
        return HttpResponse("No resumes found matching the selection.", status=404)

    # Create ZIP in-memory
    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for s in qs:
            if s.resume:
                try:
                    file_name = os.path.basename(s.resume.name)
                    first_name = s.user.first_name or ''
                    last_name = s.user.last_name or ''
                    safe_name = f"{s.usn}_{first_name}_{last_name}_{file_name}".replace(' ', '_').replace('/', '_')
                    
                    with s.resume.open('rb') as f:
                        zip_file.writestr(safe_name, f.read())
                except Exception:
                    continue

    zip_buffer.seek(0)
    response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = 'attachment; filename="student_resumes.zip"'
    return response


############################## ADMIN TRAINER VIEWS ##############################

@login_required
@user_passes_test(is_admin)
def admin_trainer_list(request):
    q = request.GET.get("q", "").strip()
    page = int(request.GET.get("page", 1))
    per_page = int(request.GET.get("per_page", 20))

    # Order the queryset to remove UnorderedObjectListWarning
    trainers = (
        TrainerProfile.objects
        .select_related("user")
        .prefetch_related("trainer_skills__domain", "trainer_skills__subdomain")
        .order_by("user__first_name", "user__last_name", "user__email")
    )

    if q:
        trainers = trainers.filter(
            Q(user__first_name__icontains=q) |
            Q(user__last_name__icontains=q) |
            Q(user__email__icontains=q) |
            Q(skill_summary__icontains=q) |
            Q(trainer_skills__domain__domain_name__icontains=q) |
            Q(trainer_skills__subdomain__subdomain_name__icontains=q)
        ).distinct()

    paginator = Paginator(trainers, per_page)
    page_obj = paginator.get_page(page)

    results = []
    for t in page_obj:
        # build full name safely even if last_name is blank
        first = (t.user.first_name or "").strip()
        last = (t.user.last_name or "").strip()
        full_name = f"{first} {last}".strip() or t.user.username

        # flatten unique domains/subdomains from prefetched trainer_skills
        doms = set()
        subs = set()
        for ts in t.trainer_skills.all():
            if getattr(ts.domain, "domain_name", None):
                doms.add(ts.domain.domain_name)
            if ts.subdomain and getattr(ts.subdomain, "subdomain_name", None):
                subs.add(ts.subdomain.subdomain_name)

        results.append({
            "id": str(t.user_id),  # primary key of TrainerProfile is user_id
            "name": full_name,
            "email": t.user.email,
            "skill_summary": t.skill_summary,
            "domains": sorted(doms),
            "subdomains": sorted(subs),
            "profile_pdf_url": (t.profile_pdf.url if t.profile_pdf else None),
        })

    return JsonResponse({
        "results": results,
        "page": page_obj.number,
        "num_pages": paginator.num_pages,
        "total": paginator.count,
        "has_next": page_obj.has_next(),
        "has_prev": page_obj.has_previous(),
    })
# ---------------------------
# Delete a trainer
# ---------------------------
@login_required
@user_passes_test(is_admin)
@require_http_methods(["DELETE"])
def admin_trainer_delete(request, trainer_id):
    trainer = get_object_or_404(TrainerProfile, pk=trainer_id)
    user = trainer.user
    user.delete()
    return JsonResponse({"success": True})


# ---------------------------
# Download trainer PDF
# ---------------------------
@login_required
@user_passes_test(is_admin)
def admin_trainer_download_profile(request, trainer_id):
    trainer = get_object_or_404(TrainerProfile, pk=trainer_id)
    if not trainer.profile_pdf:
        raise Http404("No profile PDF found")

    response = HttpResponse(trainer.profile_pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{trainer.user.full_name}_profile.pdf"'
    return response







# --------------------------- # ADMIN ATTENDANCE VIEWS # --------------------------- #

@login_required
@user_passes_test(is_admin)
def admin_attendance_list(request):
    """
    Returns attendance records + per-student summary.
    Filters: college, course, year, semester, section, date range
    """
    start = request.GET.get("start")
    end = request.GET.get("end")

    qs = AttendanceRecord.objects.select_related(
        "session", "student", "student__user",
        "session__college", "session__course", "session__section"   # OK fixed lowercase
    )

    # Apply filters on session/student
    college_raw = request.GET.getlist('college') or request.GET.get('college')
    if college_raw:
        college_ids = []
        if isinstance(college_raw, list):
            for val in college_raw:
                college_ids.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            college_ids = [v.strip() for v in college_raw.split(',') if v.strip()]
        if college_ids:
            qs = qs.filter(session__college_id__in=college_ids)

    course_raw = request.GET.getlist('course') or request.GET.get('course')
    if course_raw:
        course_ids = []
        if isinstance(course_raw, list):
            for val in course_raw:
                course_ids.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            course_ids = [v.strip() for v in course_raw.split(',') if v.strip()]
        if course_ids:
            qs = qs.filter(session__course_id__in=course_ids)

    year_raw = request.GET.getlist('year') or request.GET.get('year')
    if year_raw:
        year_list = []
        if isinstance(year_raw, list):
            for val in year_raw:
                year_list.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            year_list = [v.strip() for v in year_raw.split(',') if v.strip()]
        if year_list:
            qs = qs.filter(student__year__in=year_list)

    semester_raw = request.GET.getlist('semester') or request.GET.get('semester')
    if semester_raw:
        sem_list = []
        if isinstance(semester_raw, list):
            for val in semester_raw:
                sem_list.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            sem_list = [v.strip() for v in semester_raw.split(',') if v.strip()]
        if sem_list:
            qs = qs.filter(session__semester__in=sem_list)

    section_raw = request.GET.getlist('section') or request.GET.get('section')
    if section_raw:
        section_ids = []
        if isinstance(section_raw, list):
            for val in section_raw:
                section_ids.extend([v.strip() for v in val.split(',') if v.strip()])
        else:
            section_ids = [v.strip() for v in section_raw.split(',') if v.strip()]
        if section_ids:
            qs = qs.filter(session__section_id__in=section_ids)

    if start:
        qs = qs.filter(session__date__gte=start)
    if end:
        qs = qs.filter(session__date__lte=end)

    records = []
    per_student = {}
    present_count = 0
    absent_count = 0

    for rec in qs:
        student = rec.student
        session = rec.session

        record = {
            "student_usn": student.usn,
            "student_name": f"{student.user.first_name} {student.user.last_name}".strip(),
            "college": session.college.name if session.college else "",
            "course": session.course.name if session.course else "",
            "year": getattr(student, "year", ""),
            "semester": session.semester,
            "section": getattr(session.section, "name", ""),  # OK fixed lowercase
            "date": session.date.strftime("%Y-%m-%d"),
            "slot": session.slot_no,
            "status": rec.status,
        }
        records.append(record)

        if rec.status == "Present":
            present_count += 1
        else:
            absent_count += 1

        sid = student.usn
        if sid not in per_student:
            per_student[sid] = {
                "student_usn": student.usn,
                "student_name": f"{student.user.first_name} {student.user.last_name}".strip(),
                "present": 0,
                "absent": 0,
            }
        if rec.status == "Present":
            per_student[sid]["present"] += 1
        else:
            per_student[sid]["absent"] += 1

    per_student_list = []
    for sid, sdata in per_student.items():
        total = sdata["present"] + sdata["absent"]
        percentage = (sdata["present"] / total * 100.0) if total > 0 else 0.0
        sdata["attendance_percentage"] = round(percentage, 2)
        per_student_list.append(sdata)

    summary = {
        "total_records": len(records),
        "present": present_count,
        "absent": absent_count,
        "attendance_percentage": round((present_count / len(records) * 100.0), 2) if records else 0.0,
    }

    return JsonResponse({
        "records": records,
        "per_student": per_student_list,
        "summary": summary,
    })


@login_required
@user_passes_test(is_admin)
def admin_attendance_download(request):
    """
    Download attendance Excel file with:
    - Corner logo
    - Green styled header: 'STUDENT ATTENDANCE'
    - Attendance data table starting from row 6
    """

    import json, openpyxl
    from io import BytesIO
    from datetime import datetime
    from django.http import HttpResponse
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Fetch attendance data
    response = admin_attendance_list(request)
    if response.status_code != 200:
        return response

    data = json.loads(response.content.decode("utf-8"))

    # ---------- Workbook setup ----------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    green_fill = PatternFill(start_color="008037", end_color="008037", fill_type="solid")

    # ---------- Logo ----------
    logo_path = "static/images/atomshaalelogo.png"  # uploaded logo path
    try:
        img = Image(logo_path)
        img.width, img.height = 160, 70
        ws.add_image(img, "A1")
    except Exception as e:
        print("Logo load failed:", e)

    # ---------- Header Banner ----------
    ws.merge_cells("C1:O3")
    ws["C1"] = "STUDENT ATTENDANCE"
    ws["C1"].font = Font(size=22, bold=True, color="FFFFFF", name="Calibri")
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")
    for row in ws["C1:O3"]:
        for cell in row:
            cell.fill = green_fill

    # ---------- Table Header (starts at row 6) ----------
    start_row = 6
    headers = [
        "USN", "Name", "College", "Course", "Year",
        "Semester", "Section", "Date", "Slot", "Status"
    ]
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = green_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=Side(style="thin"), right=Side(style="thin"),
                             top=Side(style="thin"), bottom=Side(style="thin"))

    # ---------- Table Data (from row 7 onward) ----------
    data_start = start_row + 1
    for i, row in enumerate(data["records"], start=data_start):
        ws.cell(row=i, column=1, value=row["student_usn"])
        ws.cell(row=i, column=2, value=row["student_name"])
        ws.cell(row=i, column=3, value=row["college"])
        ws.cell(row=i, column=4, value=row["course"])
        ws.cell(row=i, column=5, value=row["year"])
        ws.cell(row=i, column=6, value=row["semester"])
        ws.cell(row=i, column=7, value=row["section"])
        ws.cell(row=i, column=8, value=row["date"])
        ws.cell(row=i, column=9, value=row["slot"])
        ws.cell(row=i, column=10, value=row["status"])

    # ---------- Auto-fit Columns ----------
    for i, column_cells in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_length + 3, 12), 40)

    # ---------- Alignment & Borders ----------
    for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=Side(style="thin"), right=Side(style="thin"),
                                 top=Side(style="thin"), bottom=Side(style="thin"))

    # ---------- Footer ----------
    footer_row = ws.max_row + 3
    ws.merge_cells(f"A{footer_row}:J{footer_row}")
    footer = ws[f"A{footer_row}"]
    footer.value = "Atomshaale — Powered by ATOM"
    footer.font = Font(italic=True, color="808080", size=10)
    footer.alignment = Alignment(horizontal="center", vertical="center")

    # ---------- Export ----------
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"attendance_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    resp = HttpResponse(
        buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp

# ---------------------------# Admin Feedback Views # --------------------------- #

def _feedback_queryset(request):
    feedback_qs = SessionTrainerFeedback.objects.select_related(
        "session",
        "session__college",
        "session__course",
        "session__section",
        "trainer",
        "trainer__user",
        "feedback",
        "feedback__student",
        "feedback__student__college",
        "feedback__student__course",
        "feedback__student__section",
    )

    college_ids = request.GET.get("college")
    course_ids = request.GET.get("course")
    years = request.GET.get("year")
    semesters = request.GET.get("semester")
    section_names = request.GET.get("section")
    trainer_ids = request.GET.get("trainer")
    start = parse_date(request.GET.get("start") or "")
    end = parse_date(request.GET.get("end") or "")

    def get_list(param):
        if not param or param == "all":
            return []
        return [p.strip() for p in param.split(",") if p.strip()]

    c_list = get_list(college_ids)
    if c_list:
        feedback_qs = feedback_qs.filter(session__college_id__in=c_list)

    co_list = get_list(course_ids)
    if co_list:
        feedback_qs = feedback_qs.filter(session__course_id__in=co_list)

    yr_list = get_list(years)
    if yr_list:
        try:
            yr_ints = [int(y) for y in yr_list]
            feedback_qs = feedback_qs.filter(session__year__in=yr_ints)
        except ValueError:
            pass

    sem_list = get_list(semesters)
    if sem_list:
        try:
            sem_ints = [int(s) for s in sem_list]
            feedback_qs = feedback_qs.filter(session__semester__in=sem_ints)
        except ValueError:
            pass

    s_list = get_list(section_names)
    if s_list:
        feedback_qs = feedback_qs.filter(session__section__name__in=s_list)

    t_list = get_list(trainer_ids)
    if t_list:
        feedback_qs = feedback_qs.filter(trainer_id__in=t_list)

    if start:
        feedback_qs = feedback_qs.filter(session__date__gte=start)
    if end:
        feedback_qs = feedback_qs.filter(session__date__lte=end)

    return feedback_qs


def _feedback_rows(feedback_qs):
    grouped = (
        feedback_qs
        .values(
            "session_id",
            "session__college_id",
            "session__college__name",
            "session__course_id",
            "session__course__name",
            "session__year",
            "session__semester",
            "session__section_id",
            "session__section__name",
            "trainer_id",
            "trainer__user__first_name",
            "trainer__user__last_name",
            "session__date",
            "session__slot_no",
        )
        .annotate(feedback_count=Count("id"))
        .order_by("-session__date", "session__slot_no")
    )

    rows = []
    for row in grouped:
        trainer_name = (
            f"{row.get('trainer__user__first_name', '')} "
            f"{row.get('trainer__user__last_name', '')}"
        ).strip()

        rows.append({
            "session_id": str(row.get("session_id") or ""),
            "college": row.get("session__college__name") or "",
            "course": row.get("session__course__name") or "",
            "year": row.get("session__year") or "",
            "semester": row.get("session__semester") or "",
            "section": row.get("session__section__name") or "",
            "trainer": trainer_name or "-",
            "date": row.get("session__date").strftime("%Y-%m-%d") if row.get("session__date") else "",
            "slot": row.get("session__slot_no") or "",
            "feedback_count": row.get("feedback_count") or 0,
            "college_id": str(row.get("session__college_id") or ""),
            "course_id": str(row.get("session__course_id") or ""),
            "section_id": str(row.get("session__section_id") or ""),
            "trainer_id": str(row.get("trainer_id") or ""),
        })

    return rows


def _feedback_summary(feedback_qs):
    return {
        "avg_content_rating": round(
            feedback_qs.aggregate(avg=Avg("feedback__content_rating"))["avg"] or 0,
            2,
        ),
        "avg_satisfaction": round(
            feedback_qs.aggregate(avg=Avg("feedback__satisfaction"))["avg"] or 0,
            2,
        ),
        "avg_trainer_knowledge": round(
            feedback_qs.aggregate(avg=Avg("knowledge_rating"))["avg"] or 0,
            2,
        ),
        "total_feedback": feedback_qs.count(),
    }

@login_required
@user_passes_test(is_admin)
def get_sections(request):
    """Return all available sections for dropdown."""
    try:
        course_id = request.GET.get("course")
        sections = Section.objects.all()
        if course_id:
            sections = sections.filter(sessionschedule__course_id=course_id).distinct()
        data = [{"id": s.name, "name": s.name} for s in sections.order_by("name")]
        return JsonResponse(data, safe=False)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

@login_required
@user_passes_test(is_admin)
def admin_feedback_list(request):
    feedback_qs = _feedback_queryset(request)
    return JsonResponse({
        "rows": _feedback_rows(feedback_qs),
        "summary": _feedback_summary(feedback_qs),
    })


@login_required
@user_passes_test(is_admin)
def admin_feedback_session_detail(request):
    trainer = request.GET.get("trainer")
    date = request.GET.get("date")
    slot = request.GET.get("slot")
    section = request.GET.get("section")
    semester = request.GET.get("semester")
    year = request.GET.get("year")
    course = request.GET.get("course")
    college = request.GET.get("college")

    if not all([trainer, date, slot, section, semester, year, course, college]):
        raise Http404("Missing session filters")

    feedback_qs = (
        SessionTrainerFeedback.objects
        .select_related(
            "session",
            "session__college",
            "session__course",
            "session__section",
            "trainer",
            "trainer__user",
            "feedback",
            "feedback__student",
            "feedback__student__user",
        )
        .filter(
            trainer_id=trainer,
            session__date=parse_date(date) or date,
            session__slot_no=slot,
            session__section_id=section,
            session__semester=semester,
            session__year=year,
            session__course_id=course,
            session__college_id=college,
        )
        .order_by("feedback__student__usn")
    )

    if not feedback_qs.exists():
        raise Http404("No feedback found for this session")

    first = feedback_qs.first()
    students = []
    for fb in feedback_qs:
        student = fb.feedback.student
        students.append({
            "feedback_id": str(fb.id),
            "usn": student.usn,
            "student_name": _student_name(student),
            "student_college": student.college.name if student.college else "",
            "student_course": student.course.name if student.course else "",
            "student_year": student.year,
            "student_semester": student.semester,
            "student_section": student.section.name if student.section else "",
            "trainer_name": fb.trainer.user.get_full_name() if fb.trainer and fb.trainer.user else "-",
            "session_date": fb.session.date,
            "session_slot": fb.session.slot_no,
            "content": fb.feedback.content_rating,
            "clarity": fb.feedback.objectives_clarity,
            "structure": fb.feedback.structure_logic,
            "satisfaction": fb.feedback.satisfaction,
            "engagement": fb.feedback.session_engagement,
            "future_interest": fb.feedback.future_interest,
            "takeaways": fb.feedback.takeaways,
            "pace_rating": fb.pace_rating,
            "knowledge": fb.knowledge_rating,
            "queries": fb.queries_rating,
            "participation": fb.participation,
            "examples": fb.examples,
        })

    context = {
        "session_info": {
            "college": first.session.college.name,
            "course": first.session.course.name,
            "year": first.session.year,
            "semester": first.session.semester,
            "section": first.session.section.name,
            "trainer": first.session.trainer.user.get_full_name() if first.session.trainer and first.session.trainer.user else "-",
            "date": first.session.date,
            "slot": first.session.slot_no,
            "total_feedbacks": feedback_qs.count(),
        },
        "students": students,
        "feedbacks": students,
    }

    return render(request, "admin_feedback_session_detail.html", context)

@login_required
@user_passes_test(is_admin)
def admin_feedback_detail(request, feedback_id):

    feedback = get_object_or_404(
        SessionTrainerFeedback.objects.select_related(
            "feedback",
            "feedback__student",
            "feedback__student__user",
            "feedback__student__college",
            "feedback__student__course",
            "trainer",
            "trainer__user",
            "session"
        ),
        pk=feedback_id
    )

    sf = feedback.feedback
    student = sf.student

    context = {

        "feedback": feedback,

        "student": student,

        "trainer": feedback.trainer,

        "session": feedback.session,
    }

    return render(
        request,
        "admin_feedback_detail.html",
        context
    )

@login_required
@user_passes_test(is_admin)
def admin_feedback_download(request):
    """Download feedback as CSV"""
    dt = localtime().strftime("%Y%m%d_%H%M%S")
    filename = f"feedback_{dt}.csv"

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)

    writer.writerow([
        "Trainer",
        "Student USN",
        "Student Name",
        "College",
        "Course",
        "Year",
        "Semester",
        "Section",
        "Date",
        "Slot",
        "Content Rating",
        "Objectives Clarity",
        "Structure Logic",
        "Satisfaction",
        "Engagement",
        "Future Interest",
        "Takeaways",
        "Pace Rating",
        "Knowledge Rating",
        "Queries Rating",
        "Participation",
        "Examples",
    ])

    for r in _feedback_queryset(request):
        student = r.feedback.student
        session = r.session
        writer.writerow([
            r.trainer.user.get_full_name() if getattr(r.trainer, "user", None) else str(r.trainer),
            student.usn,
            student.user.get_full_name(),
            student.college.name,
            student.course.name if student.course else "",
            student.year,
            student.semester,
            student.section.name if student.section else "",
            session.date.strftime("%Y-%m-%d") if session and session.date else "",
            session.slot_no if session else "",
            r.feedback.content_rating,
            r.feedback.objectives_clarity,
            r.feedback.structure_logic,
            r.feedback.satisfaction,
            r.feedback.session_engagement,
            r.feedback.future_interest,
            r.feedback.takeaways,
            r.pace_rating,
            r.knowledge_rating,
            r.queries_rating,
            r.participation,
            r.examples,
        ])

    return response

# --------------------------- # Admin session management   # --------------------------- #
@require_GET
@login_required
@user_passes_test(is_admin)
def get_colleges(request):
    colleges = College.objects.values("id", "name")
    return JsonResponse(list(colleges), safe=False)


@require_GET
@login_required
@user_passes_test(is_admin)
def get_courses(request):
    """
    Return all available courses (since Course is independent).
    """
    try:
        from college.models import Course
        courses = Course.objects.values("id", "name")
        data = [{"id": c["id"], "name": c["name"]} for c in courses]
        return JsonResponse(data, safe=False)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


@require_GET
@login_required
@user_passes_test(is_admin)
def get_semester_choices(request):
    """Fetch available semester choices directly from StudentProfile model."""
    try:
        semester_field = StudentProfile._meta.get_field("semester")
        semesters = [{"id": value, "label": label} for value, label in semester_field.choices]
        return JsonResponse(semesters, safe=False)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


@require_GET
@login_required
@user_passes_test(is_admin)
def get_year_choices(request):
    """Fetch available year choices directly from StudentProfile model."""
    try:
        year_field = StudentProfile._meta.get_field("year")
        years = [{"id": value, "label": label} for value, label in year_field.choices]
        return JsonResponse(years, safe=False)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


@require_GET
@login_required
@user_passes_test(is_admin)
def get_sections(request):
    """Return available sections for the selected college and course(s)."""
    try:
        college_id = request.GET.get("college")
        course_raw = request.GET.get("course", "")
        course_ids = [c.strip() for c in course_raw.split(",") if c.strip()]
        
        sections = Section.objects.all()
        
        if college_id and course_ids:
            from college.models import CollegeCourseSection
            sections = sections.filter(
                college_course_sections__college_course__college_id=college_id,
                college_course_sections__college_course__course_id__in=course_ids
            ).distinct()
        elif college_id:
            from college.models import CollegeCourseSection
            sections = sections.filter(
                college_course_sections__college_course__college_id=college_id
            ).distinct()
        elif course_ids:
            from college.models import CollegeCourseSection
            sections = sections.filter(
                college_course_sections__college_course__course_id__in=course_ids
            ).distinct()
        elif course_raw:
            sections = sections.filter(sessionschedule__course_id=course_raw).distinct()
            
        data = [{"id": s.name, "name": s.name} for s in sections.order_by("name")]
        return JsonResponse(data, safe=False)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

@require_GET
@login_required
@user_passes_test(is_admin)
def get_domains(request):
    domains = Domain.objects.values("id", "domain_name")
    return JsonResponse(list(domains), safe=False)


@require_GET
@login_required
@user_passes_test(is_admin)
def get_subdomains(request, domain_id):
    """Return all subdomains for the selected domain."""
    try:
        subdomains = SubDomain.objects.filter(domain__id=domain_id).values("id", "subdomain_name")
        return JsonResponse(list(subdomains), safe=False)
    except SubDomain.DoesNotExist:
        return JsonResponse([], safe=False)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

@require_GET
@login_required
@user_passes_test(is_admin)
def get_trainers(request):
    """
    Return a list of all trainers for dropdown population.
    Uses 'user_id' as primary key since TrainerProfile has no 'id' field.
    """
    trainers = TrainerProfile.objects.values("user_id", "user__first_name", "user__last_name")
    data = [
        {
            "id": t["user_id"],
            "name": f"{t['user__first_name']} {t['user__last_name']}"
        }
        for t in trainers
    ]
    return JsonResponse(data, safe=False)


# ------------------- CREATE SESSION -------------------
@require_POST
@login_required
@user_passes_test(is_admin)
def create_session_schedule(request):
    """
    Create a new session schedule entry.
    Handles validation and logs clear error messages for debugging.
    """
    try:
        form_data = {
            "college_id": request.POST.get("college"),
            "course_id": request.POST.get("course"),
            "section_id": request.POST.get("section"),
            "domain_id": request.POST.get("domain"),
            "subdomain_id": request.POST.get("subdomain"),
            "module_id": request.POST.get("module"),
            "trainer_id": request.POST.get("trainer"),
            "semester": request.POST.get("semester"),
            "year": request.POST.get("year"),
            "date": request.POST.get("date"),       # OK Correct field name
            "slot_no": request.POST.get("slot_no"), # OK Match HTML <select name="slot_no">
        }

        # Clean empty or missing values
        safe_data = {k: v for k, v in form_data.items() if v not in [None, ""]}

        # Validate required fields
        required_fields = ["date", "slot_no", "college_id", "course_id", "section_id", "semester", "year"]
        missing = [f for f in required_fields if f not in safe_data]
        if missing:
            return JsonResponse({
                "success": False,
                "message": f"WARNING Missing required fields: {', '.join(missing)}"
            }, status=400)

        # OK Create session schedule
        schedule = SessionSchedule.objects.create(**safe_data)

        return JsonResponse({
            "success": True,
            "message": f"OK Session scheduled successfully for {schedule.date} (Slot {schedule.slot_no})."
        })

    except Exception as e:
        print("ERROR Error creating schedule:", e)
        return JsonResponse({
            "success": False,
            "message": f"Error creating schedule: {e}"
        }, status=500)

@require_GET
@login_required
@user_passes_test(is_admin)
def get_modules(request):
    """
    Return a list of modules filtered by domain and subdomain.
    """
    try:
        domain_id = request.GET.get("domain_id")
        subdomain_id = request.GET.get("subdomain_id")
        
        filters = {}
        if domain_id:
            filters["domain_id"] = domain_id
        if subdomain_id:
            filters["subdomain_id"] = subdomain_id
            
        modules = Module.objects.filter(**filters).order_by("module_name")
        data = [{"id": str(m.id), "name": m.module_name} for m in modules]
        return JsonResponse(data, safe=False)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

@require_POST
@login_required
@user_passes_test(is_admin)
def bulk_create_session_schedule(request):
    """
    Create multiple training session schedules at once.
    """
    try:
        data = json.loads(request.body)
        schedules_data = data.get("schedules", [])
        if not schedules_data:
            return JsonResponse({"success": False, "message": "No schedules provided."}, status=400)
        
        created_schedules = []
        errors = []
        
        with transaction.atomic():
            for index, item in enumerate(schedules_data):
                try:
                    college_id = item.get("college")
                    course_id = item.get("course")
                    section_id = item.get("section")
                    domain_id = item.get("domain")
                    subdomain_id = item.get("subdomain") or None
                    module_id = item.get("module") or None
                    trainer_id = item.get("trainer") or None
                    semester = item.get("semester")
                    year = item.get("year")
                    date = item.get("date")
                    slot_no = item.get("slot_no")
                    
                    if not (college_id and course_id and section_id and semester and year and date and slot_no and domain_id):
                        raise ValidationError(f"Row {index+1}: Missing required fields.")
                    
                    # Convert types
                    try:
                        slot_no = int(slot_no)
                        semester = int(semester)
                        year = int(year)
                    except ValueError:
                        raise ValidationError(f"Row {index+1}: Slot, Semester, and Year must be integers.")
                    
                    # Trainer availability conflict check
                    if trainer_id:
                        trainer_conflict = SessionSchedule.objects.filter(
                            date=date, slot_no=slot_no, trainer_id=trainer_id
                        ).exists()
                        if trainer_conflict:
                            tp = TrainerProfile.objects.get(user_id=trainer_id)
                            raise ValidationError(f"Row {index+1}: Trainer {tp.user.get_full_name()} is already scheduled for Slot {slot_no} on {date}.")
                    
                    # Session schedule uniqueness conflict check
                    session_conflict = SessionSchedule.objects.filter(
                        date=date, slot_no=slot_no, section_id=section_id,
                        course_id=course_id, semester=semester, college_id=college_id
                    ).exists()
                    if session_conflict:
                        raise ValidationError(f"Row {index+1}: A session is already scheduled for Slot {slot_no} in this section on {date}.")
                    
                    schedule = SessionSchedule.objects.create(
                        college_id=college_id,
                        course_id=course_id,
                        section_id=section_id,
                        domain_id=domain_id,
                        subdomain_id=subdomain_id,
                        module_id=module_id,
                        trainer_id=trainer_id,
                        semester=semester,
                        year=year,
                        date=date,
                        slot_no=slot_no,
                    )
                    created_schedules.append(schedule)
                except ValidationError as ve:
                    errors.append(str(ve))
                except Exception as ex:
                    errors.append(f"Row {index+1} error: {ex}")
            
            if errors:
                raise ValidationError("; ".join(errors))
                
        return JsonResponse({"success": True, "message": f"OK Successfully scheduled {len(created_schedules)} sessions."})
    except ValidationError as ve:
        return JsonResponse({"success": False, "message": str(ve)}, status=400)
    except Exception as e:
        return JsonResponse({"success": False, "message": f"Server error: {e}"}, status=500)

@require_POST
@login_required
@user_passes_test(is_admin)
def bulk_upload_session_schedule_excel(request):
    """
    Parse uploaded Excel sheet and schedule training sessions.
    """
    excel_file = request.FILES.get("file")
    if not excel_file:
        return JsonResponse({"success": False, "message": "No file uploaded."}, status=400)
    
    try:
        filename = excel_file.name.lower()
        if filename.endswith('.csv'):
            try:
                csv_data = excel_file.read().decode('utf-8-sig').splitlines()
                reader = csv.reader(csv_data)
                rows = list(reader)
            except Exception as e:
                return JsonResponse({"success": False, "message": f"Failed to parse CSV file: {e}"}, status=400)
        else:
            try:
                wb = openpyxl.load_workbook(excel_file, read_only=True)
                sheet = wb.active
                rows = list(sheet.iter_rows(values_only=True))
            except Exception as e:
                return JsonResponse({"success": False, "message": f"Failed to parse Excel file: {e}. If this is a CSV file, please ensure it has a .csv extension."}, status=400)
        
        if len(rows) < 2:
            return JsonResponse({"success": False, "message": "The uploaded file has no data rows."}, status=400)
        
        headers = [str(h).strip().lower() for h in rows[0] if h is not None]
        required_cols = ["date", "slot", "college", "course", "section", "year", "semester", "domain"]
        missing = [c for c in required_cols if c not in headers]
        if missing:
            return JsonResponse({
                "success": False,
                "message": f"Missing required columns in sheet header: {', '.join(missing)}"
            }, status=400)
            
        col_map = {h: i for i, h in enumerate(headers)}
        errors = []
        created_count = 0
        
        with transaction.atomic():
            for row_idx, row in enumerate(rows[1:], start=2):
                if not any(row):  # skip empty lines
                    continue
                
                try:
                    def get_val(col_name, required=True):
                        idx = col_map.get(col_name)
                        if idx is None or idx >= len(row):
                            if required:
                                raise ValidationError(f"Column '{col_name}' is missing.")
                            return None
                        val = row[idx]
                        if val is None or str(val).strip() == "":
                            if required:
                                raise ValidationError(f"Column '{col_name}' is required.")
                            return None
                        return val
                    
                    # Parse Date
                    date_val = get_val("date")
                    import datetime as dt_module
                    if isinstance(date_val, dt_module.datetime):
                        date_parsed = date_val.date()
                    elif isinstance(date_val, dt_module.date):
                        date_parsed = date_val
                    else:
                        date_str = str(date_val).strip()
                        try:
                            date_parsed = datetime.strptime(date_str, "%Y-%m-%d").date()
                        except ValueError:
                            try:
                                date_parsed = datetime.strptime(date_str, "%d-%m-%Y").date()
                            except ValueError:
                                raise ValidationError(f"Invalid date format '{date_str}'. Use YYYY-MM-DD or DD-MM-YYYY.")
                    
                    # Parse Slot
                    slot_val = get_val("slot")
                    try:
                        slot_no = int(slot_val)
                        if slot_no not in [1, 2, 3, 4, 5]:
                            raise ValueError
                    except ValueError:
                        raise ValidationError(f"Slot must be an integer between 1 and 5 (got '{slot_val}').")
                    
                    # Resolve College
                    college_name = str(get_val("college")).strip()
                    try:
                        college = College.objects.get(name__iexact=college_name)
                    except College.DoesNotExist:
                        raise ValidationError(f"College '{college_name}' not found.")
                    
                    # Resolve Course
                    course_name = str(get_val("course")).strip()
                    try:
                        course = Course.objects.get(name__iexact=course_name)
                    except Course.DoesNotExist:
                        raise ValidationError(f"Course '{course_name}' not found.")
                    
                    # Resolve Section
                    section_name = str(get_val("section")).strip()
                    try:
                        section = Section.objects.get(name__iexact=section_name)
                    except Section.DoesNotExist:
                        raise ValidationError(f"Section '{section_name}' not found.")
                    
                    # Parse Year and Semester
                    try:
                        year = int(get_val("year"))
                    except ValueError:
                        raise ValidationError(f"Year must be an integer.")
                    
                    try:
                        semester = int(get_val("semester"))
                    except ValueError:
                        raise ValidationError(f"Semester must be an integer.")
                    
                    # Resolve Domain
                    domain_name = str(get_val("domain")).strip()
                    try:
                        domain = Domain.objects.get(domain_name__iexact=domain_name)
                    except Domain.DoesNotExist:
                        raise ValidationError(f"Domain '{domain_name}' not found.")
                    
                    # Resolve Subdomain
                    subdomain_name = get_val("subdomain", required=False)
                    subdomain = None
                    if subdomain_name:
                        subdomain_name = str(subdomain_name).strip()
                        try:
                            subdomain = SubDomain.objects.get(subdomain_name__iexact=subdomain_name, domain=domain)
                        except SubDomain.DoesNotExist:
                            raise ValidationError(f"Subdomain '{subdomain_name}' not found under domain '{domain_name}'.")
                    
                    # Resolve Module
                    module_name = get_val("module", required=False)
                    module = None
                    if module_name:
                        module_name = str(module_name).strip()
                        try:
                            m_filters = {"module_name__iexact": module_name, "domain": domain}
                            if subdomain:
                                m_filters["subdomain"] = subdomain
                            module = Module.objects.get(**m_filters)
                        except Module.DoesNotExist:
                            raise ValidationError(f"Module '{module_name}' not found under domain '{domain_name}' and subdomain '{subdomain_name or 'None'}'.")
                    
                    # Resolve Trainer
                    trainer_name = get_val("trainer", required=False)
                    trainer = None
                    if trainer_name:
                        trainer_name = str(trainer_name).strip()
                        parts = trainer_name.split()
                        if len(parts) == 1:
                            t_q = Q(user__first_name__iexact=parts[0]) | Q(user__last_name__iexact=parts[0]) | Q(user__email__iexact=parts[0])
                        else:
                            t_q = Q(user__first_name__iexact=parts[0], user__last_name__iexact=parts[1]) | Q(user__first_name__iexact=trainer_name)
                        try:
                            trainer = TrainerProfile.objects.get(t_q)
                        except TrainerProfile.DoesNotExist:
                            raise ValidationError(f"Trainer '{trainer_name}' not found.")
                        except TrainerProfile.MultipleObjectsReturned:
                            raise ValidationError(f"Multiple trainers match name '{trainer_name}'. Be more specific.")
                    
                    # Unique checks
                    if trainer:
                        trainer_conflict = SessionSchedule.objects.filter(
                            date=date_parsed, slot_no=slot_no, trainer=trainer
                        ).exists()
                        if trainer_conflict:
                            raise ValidationError(f"Trainer '{trainer.user.get_full_name()}' is already scheduled for Slot {slot_no} on {date_parsed}.")
                    
                    session_conflict = SessionSchedule.objects.filter(
                        date=date_parsed, slot_no=slot_no, section=section,
                        course=course, semester=semester, college=college
                    ).exists()
                    if session_conflict:
                        raise ValidationError(f"Slot {slot_no} in section '{section_name}' is already scheduled on {date_parsed}.")
                    
                    SessionSchedule.objects.create(
                        date=date_parsed,
                        slot_no=slot_no,
                        college=college,
                        course=course,
                        section=section,
                        year=year,
                        semester=semester,
                        domain=domain,
                        subdomain=subdomain,
                        module=module,
                        trainer=trainer
                    )
                    created_count += 1
                except ValidationError as ve:
                    errors.append(f"Row {row_idx}: {ve.message}")
                except Exception as ex:
                    errors.append(f"Row {row_idx}: {ex}")
            
            if errors:
                raise ValidationError("; ".join(errors))
                
        return JsonResponse({"success": True, "message": f"OK Successfully imported {created_count} schedules."})
    except ValidationError as ve:
        return JsonResponse({"success": False, "message": str(ve)}, status=400)
    except Exception as e:
        return JsonResponse({"success": False, "message": f"Error parsing Excel: {e}"}, status=500)

@require_GET
@login_required
@user_passes_test(is_admin)
def list_session_schedules(request):
    """
    Returns all session schedules in JSON format for the admin panel.
    """
    try:
        schedules = SessionSchedule.objects.select_related(
            "college", "course", "section", "domain", "subdomain", "trainer", "module"
        ).order_by("-date", "slot_no")

        data = []
        for s in schedules:
            data.append({
                "id": str(s.id),
                "date": s.date.strftime("%Y-%m-%d"),
                "slot_no": s.slot_no,
                "college": s.college.name if s.college else None,
                "college_id": str(s.college_id) if s.college_id else None,
                "course": s.course.name if s.course else None,
                "course_id": str(s.course_id) if s.course_id else None,
                "section": s.section.name if s.section else None,
                "section_id": s.section.name if s.section else None,
                "domain": s.domain.domain_name if s.domain else None,
                "domain_id": str(s.domain_id) if s.domain_id else None,
                "subdomain": s.subdomain.subdomain_name if s.subdomain else None,
                "subdomain_id": str(s.subdomain_id) if s.subdomain_id else None,
                "module": s.module.module_name if s.module else None,
                "module_id": str(s.module_id) if s.module_id else None,
                "trainer": s.trainer.user.get_full_name() if s.trainer else "Unassigned",
                "trainer_id": str(s.trainer_id) if s.trainer_id else None,
                "semester": s.semester,
                "year": s.year,
                "done": s.done,
                "created_at": s.created_at.strftime("%Y-%m-%d %H:%M"),
            })

        return JsonResponse({"success": True, "schedules": data})

    except Exception as e:
        print("ERROR Error fetching schedules:", e)
        return JsonResponse({
            "success": False,
            "message": f"Error fetching schedules: {e}"
        }, status=500)
    
@require_http_methods(["DELETE"])
@login_required
@user_passes_test(is_admin)
def delete_session_schedule(request, schedule_id):
    """
    Delete a specific session schedule by ID.
    """
    try:
        schedule = SessionSchedule.objects.get(id=schedule_id)
        schedule.delete()
        return JsonResponse({"success": True, "message": "OK Session deleted successfully."})
    except SessionSchedule.DoesNotExist:
        return JsonResponse({"success": False, "message": "Session not found."}, status=404)
    except Exception as e:
        print("ERROR Delete error:", e)
        return JsonResponse({"success": False, "message": str(e)}, status=500)


@require_http_methods(["POST"])
@login_required
@user_passes_test(is_admin)
def update_session_schedule(request, schedule_id):
    """
    Update specific session schedule fields.
    """
    try:
        schedule = SessionSchedule.objects.get(id=schedule_id)

        # Only update provided fields
        updatable_fields = [
            "date", "slot_no", "college_id", "course_id",
            "section_id", "domain_id", "subdomain_id", "module_id", "trainer_id",
            "semester", "year"
        ]

        for field in updatable_fields:
            if field in request.POST:
                val = request.POST[field]
                setattr(schedule, field, val if val != "" else None)

        schedule.save()

        return JsonResponse({
            "success": True,
            "message": "OK Session updated successfully.",
            "data": model_to_dict(schedule),
        })

    except SessionSchedule.DoesNotExist:
        return JsonResponse({"success": False, "message": "Session not found."}, status=404)
    except Exception as e:
        print("ERROR Update error:", e)
        return JsonResponse({"success": False, "message": str(e)}, status=500)
    

# --- API: Fetch all trainer reports ---

@login_required
@user_passes_test(is_admin)
def admin_trainer_reports_api(request):
    """Return filtered trainer session reports as JSON for admin view."""
    from_date = request.GET.get("from")
    to_date = request.GET.get("to")
    trainer_ids = request.GET.get("trainer")
    college_ids = request.GET.get("college")
    course_ids = request.GET.get("course")
    section_names = request.GET.get("section")
    domain_ids = request.GET.get("domain")
    subdomain_ids = request.GET.get("subdomain")
    years = request.GET.get("year")
    semesters = request.GET.get("semester")
    search_query = request.GET.get("search")

    reports = (
        TrainerSessionReport.objects
        .select_related(
            "trainer__user",
            "college",
            "course",
            "domain",
            "subdomain",
            "section"
        )
    )

    filters = Q()

    if from_date and to_date:
        filters &= Q(date__range=[from_date, to_date])
    elif from_date:
        filters &= Q(date__gte=from_date)
    elif to_date:
        filters &= Q(date__lte=to_date)

    def get_list(param):
        if not param or param == "all":
            return []
        return [p.strip() for p in param.split(",") if p.strip()]

    t_list = get_list(trainer_ids)
    if t_list:
        filters &= Q(trainer__user__id__in=t_list)

    c_list = get_list(college_ids)
    if c_list:
        filters &= Q(college__id__in=c_list)

    co_list = get_list(course_ids)
    if co_list:
        filters &= Q(course__id__in=co_list)

    s_list = get_list(section_names)
    if s_list:
        filters &= Q(section__name__in=s_list)

    d_list = get_list(domain_ids)
    if d_list:
        filters &= Q(domain__id__in=d_list)

    sub_list = get_list(subdomain_ids)
    if sub_list:
        filters &= Q(subdomain__id__in=sub_list)

    yr_list = get_list(years)
    if yr_list:
        try:
            yr_ints = [int(y) for y in yr_list]
            filters &= Q(year__in=yr_ints)
        except ValueError:
            pass

    sem_list = get_list(semesters)
    if sem_list:
        try:
            sem_ints = [int(s) for s in sem_list]
            filters &= Q(semester__in=sem_ints)
        except ValueError:
            pass

    if search_query:
        search_query = search_query.strip()
        filters &= (
            Q(trainer__user__first_name__icontains=search_query) |
            Q(trainer__user__last_name__icontains=search_query) |
            Q(college__name__icontains=search_query) |
            Q(course__name__icontains=search_query) |
            Q(domain__domain_name__icontains=search_query) |
            Q(subdomain__subdomain_name__icontains=search_query) |
            Q(module__icontains=search_query) |
            Q(summary__icontains=search_query) |
            Q(best_performer__icontains=search_query)
        )

    reports = reports.filter(filters).order_by("-date", "trainer__user__first_name")

    data = []
    for report in reports:
        try:
            trainer_name = report.trainer.user.get_full_name()
        except AttributeError:
            trainer_name = f"{report.trainer.user.first_name} {report.trainer.user.last_name}".strip()

        data.append({
            "id": report.id,
            "date": report.date.strftime("%Y-%m-%d"),
            "trainer": trainer_name or "Unknown Trainer",
            "college": report.college.name if report.college else "-",
            "course": report.course.name if report.course else "-",
            "year": report.year,
            "semester": report.semester,
            "section": report.section.name if report.section else "-",
            "domain": report.domain.domain_name if report.domain else "-",
            "subdomain": report.subdomain.subdomain_name if report.subdomain else "-",
            "module": report.module,
            "total_students": report.total_students,
            "present_students": report.present_students,
            "absent_students": report.absent_students,
            "exercises_solved": report.exercises_solved,
            "materials_shared": report.materials_shared,
            "assignments_given": report.assignments_given,
            "best_performer": report.best_performer or "-",
            "summary": report.summary,
            "submitted_at": report.submitted_at.strftime("%Y-%m-%d %H:%M"),
        })
    
    return JsonResponse({"reports": data})

@login_required
@user_passes_test(is_admin)
def export_trainer_reports_excel(request):
    """Export filtered or selected trainer reports to Excel."""
    from_date = request.GET.get("from")
    to_date = request.GET.get("to")
    trainer_ids = request.GET.get("trainer")
    college_ids = request.GET.get("college")
    course_ids = request.GET.get("course")
    section_names = request.GET.get("section")
    domain_ids = request.GET.get("domain")
    subdomain_ids = request.GET.get("subdomain")
    years = request.GET.get("year")
    semesters = request.GET.get("semester")
    search_query = request.GET.get("search")
    report_ids_str = request.GET.get("report_ids")

    reports = TrainerSessionReport.objects.select_related(
        "trainer__user", "college", "course", "domain", "subdomain", "section"
    )

    filters = Q()

    if report_ids_str:
        r_ids = [r.strip() for r in report_ids_str.split(",") if r.strip()]
        filters &= Q(id__in=r_ids)
    else:
        if from_date and to_date:
            filters &= Q(date__range=[from_date, to_date])
        elif from_date:
            filters &= Q(date__gte=from_date)
        elif to_date:
            filters &= Q(date__lte=to_date)

        def get_list(param):
            if not param or param == "all":
                return []
            return [p.strip() for p in param.split(",") if p.strip()]

        t_list = get_list(trainer_ids)
        if t_list:
            filters &= Q(trainer__user__id__in=t_list)

        c_list = get_list(college_ids)
        if c_list:
            filters &= Q(college__id__in=c_list)

        co_list = get_list(course_ids)
        if co_list:
            filters &= Q(course__id__in=co_list)

        s_list = get_list(section_names)
        if s_list:
            filters &= Q(section__name__in=s_list)

        d_list = get_list(domain_ids)
        if d_list:
            filters &= Q(domain__id__in=d_list)

        sub_list = get_list(subdomain_ids)
        if sub_list:
            filters &= Q(subdomain__id__in=sub_list)

        yr_list = get_list(years)
        if yr_list:
            try:
                yr_ints = [int(y) for y in yr_list]
                filters &= Q(year__in=yr_ints)
            except ValueError:
                pass

        sem_list = get_list(semesters)
        if sem_list:
            try:
                sem_ints = [int(s) for s in sem_list]
                filters &= Q(semester__in=sem_ints)
            except ValueError:
                pass

        if search_query:
            search_query = search_query.strip()
            filters &= (
                Q(trainer__user__first_name__icontains=search_query) |
                Q(trainer__user__last_name__icontains=search_query) |
                Q(college__name__icontains=search_query) |
                Q(course__name__icontains=search_query) |
                Q(domain__domain_name__icontains=search_query) |
                Q(subdomain__subdomain_name__icontains=search_query) |
                Q(module__icontains=search_query) |
                Q(summary__icontains=search_query) |
                Q(best_performer__icontains=search_query)
            )

    reports = reports.filter(filters).order_by("-date", "trainer__user__first_name")

    # ------------------ Excel setup ------------------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trainer Reports"

    headers = [
        "Date", "Trainer", "College", "Course", "Year", "Semester", "Section",
        "Domain", "Subdomain", "Module", "Total Students", "Present", "Absent",
        "Exercises Solved", "Notes Shared", "Assignments Given", "Best Performer",
        "Summary", "Submitted At"
    ]
    ws.append(headers)

    for r in reports:
        trainer_name = (
            r.trainer.user.get_full_name()
            if hasattr(r.trainer.user, "get_full_name")
            else f"{r.trainer.user.first_name} {r.trainer.user.last_name}"
        )
        ws.append([
            r.date.strftime("%Y-%m-%d"),
            trainer_name,
            r.college.name if r.college else "",
            r.course.name if r.course else "",
            r.year,
            r.semester,
            r.section.name if r.section else "",
            r.domain.domain_name if r.domain else "",
            r.subdomain.subdomain_name if r.subdomain else "",
            r.module,
            r.total_students,
            r.present_students,
            r.absent_students,
            r.exercises_solved,
            r.materials_shared,
            r.assignments_given,
            r.best_performer or "",
            r.summary,
            r.submitted_at.strftime("%Y-%m-%d %H:%M"),
        ])

    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"trainer_reports_{from_date or 'all'}_to_{to_date or 'all'}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response



@login_required
@user_passes_test(is_admin)
@require_GET
def trainer_list_api(request):
    """Return all trainers for admin dropdown."""
    trainers = TrainerProfile.objects.select_related("user").all()
    data = [
        {"id": t.user.id, "name": t.user.get_full_name()}
        for t in trainers
    ]
    return JsonResponse({"trainers": data})


@login_required
@user_passes_test(is_admin)
def admin_edit_student(request, student_id):
    student_profile = get_object_or_404(StudentProfile, id=student_id)
    user = student_profile.user
    
    if request.method == 'POST':
        form = StudentEditForm(request.POST, instance=user, student_profile=student_profile)
        if form.is_valid():
            form.save()
            messages.success(request, "Student details updated successfully.")
            next_url = request.POST.get('next') or request.GET.get('next')
            if next_url:
                return redirect(next_url)
            return redirect('college_student_page', college_id=student_profile.college.id)
    else:
        form = StudentEditForm(instance=user, student_profile=student_profile)
        
    return render(request, 'admin_edit_student.html', {
        'form': form,
        'student': student_profile,
        'next_url': request.POST.get('next') or request.GET.get('next') or request.META.get('HTTP_REFERER', '')
    })


@login_required
@user_passes_test(is_admin)
@require_POST
def admin_delete_student(request, student_id):
    student_profile = get_object_or_404(StudentProfile, id=student_id)
    user = student_profile.user
    name = user.get_full_name()
    user.delete()
    messages.success(request, f"Student {name} has been deleted.")
    
    next_url = request.POST.get('next') or request.GET.get('next') or request.META.get('HTTP_REFERER')
    if next_url:
        return redirect(next_url)
    return redirect('admin_college')


